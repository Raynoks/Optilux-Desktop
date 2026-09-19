"""
OPTILUX Desktop — application de gestion (clients, rendez-vous, stock, ventes, dépenses).
100% locale : les données sont stockées dans optilux.db (SQLite) à côté de ce fichier.
Aucune connexion internet requise pour fonctionner au quotidien.
"""
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import datetime
import os
import sys
import subprocess
import uuid
import zipfile
import shutil
from lxml import etree

from database import (init_db, get_connection, log_action, DB_PATH, get_setting, set_setting,
                       get_types, add_type, delete_type, DISCOUNT_OPTIONS, PRODUCT_IMAGES_DIR,
                       FACTURE_FILES_DIR)
from notifications import notify
from whatsapp import send_whatsapp, normalize_phone


APP_VERSION = "1.0.2"           # bump this on every release
UPDATE_URL = "https://raw.githubusercontent.com/Raynoks/Optilux-Desktop/main/latest.json"

ASSETS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets")
ICON_ICO = os.path.join(ASSETS_DIR, "optilux.ico")
LOGO_PNG = os.path.join(ASSETS_DIR, "optilux_logo.png")


def load_logo_photo(max_size):
    """Retourne une PhotoImage du logo, redimensionnée. Fonctionne avec ou sans Pillow installé."""
    if not os.path.exists(LOGO_PNG):
        return None
    try:
        from PIL import Image, ImageTk
        im = Image.open(LOGO_PNG).convert("RGBA")
        im.thumbnail((max_size, max_size), Image.LANCZOS)
        return ImageTk.PhotoImage(im)
    except Exception:
        try:
            photo = tk.PhotoImage(file=LOGO_PNG)
            factor = max(1, photo.width() // max_size)
            return photo.subsample(factor, factor)
        except Exception:
            return None


def build_pulse_frames(max_size, bg_hex, steps=16, min_opacity=0.18):
    """Pre-renders the logo faded to varying opacity over a solid background, for a smooth
    fade in/out pulse. Requires Pillow — returns None (caller falls back to static logo) if
    Pillow isn't installed."""
    if not os.path.exists(LOGO_PNG):
        return None
    try:
        from PIL import Image, ImageTk
    except Exception:
        return None
    try:
        logo = Image.open(LOGO_PNG).convert("RGBA")
        logo.thumbnail((max_size, max_size), Image.LANCZOS)
        bg = Image.new("RGBA", logo.size, bg_hex)
        frames = []
        for i in range(steps):
            t = i / (steps - 1)
            opacity = min_opacity + (1 - min_opacity) * t
            faded = logo.copy()
            r, g, b, a = faded.split()
            a = a.point(lambda px: int(px * opacity))
            faded.putalpha(a)
            composed = Image.alpha_composite(bg, faded).convert("RGB")
            frames.append(ImageTk.PhotoImage(composed))
        return frames
    except Exception:
        return None

# ---------------------------------------------------------------- THEME
INK = "#0B0B0C"
CHARCOAL = "#18181B"
LINE = "#2B2A2D"
LINE_SOFT = "#E4E1D9"
PAPER = "#F7F5F1"
RED = "#D6293B"
RED_BRIGHT = "#FF4D5E"
SLATE = "#8A8A90"
WHITE = "#FFFFFF"

# Sidebar / login / splash always stay this fixed dark "chrome", regardless of the
# light/dark toggle below — only the workspace (pages, tables, modals) re-themes.
# These content-area colors ARE reassigned by apply_theme().
THEMES = {
    "light": {
        "content_bg": "#F7F5F1",
        "surface": "#FFFFFF",
        "border": "#E4E1D9",
        "text": "#0B0B0C",
        "table_odd": "#FBFAF8",
        "table_selected": "#F6C7CD",
        "table_low": "#FCE8EA",
        "table_hover": "#FBE4E7",
        "card_dark": "#0B0B0C",
    },
    "dark": {
        "content_bg": "#111113",
        "surface": "#1C1C1F",
        "border": "#2E2E32",
        "text": "#F2F1EE",
        "table_odd": "#1A1A1D",
        "table_selected": "#5C2530",
        "table_low": "#3A1F23",
        "table_hover": "#33191C",
        "card_dark": "#242428",
    },
}
CONTENT_BG = THEMES["light"]["content_bg"]
SURFACE = THEMES["light"]["surface"]
BORDER = THEMES["light"]["border"]
TEXT = THEMES["light"]["text"]
TABLE_ODD = THEMES["light"]["table_odd"]
TABLE_SELECTED = THEMES["light"]["table_selected"]
TABLE_LOW = THEMES["light"]["table_low"]
TABLE_HOVER = THEMES["light"]["table_hover"]
CARD_DARK = THEMES["light"]["card_dark"]
THEME_NAME = "light"


def apply_theme(name):
    global CONTENT_BG, SURFACE, BORDER, TEXT, TABLE_ODD, TABLE_SELECTED, CARD_DARK, THEME_NAME
    global TABLE_LOW, TABLE_HOVER
    t = THEMES.get(name, THEMES["light"])
    CONTENT_BG, SURFACE, BORDER, TEXT = t["content_bg"], t["surface"], t["border"], t["text"]
    TABLE_ODD, TABLE_SELECTED, CARD_DARK = t["table_odd"], t["table_selected"], t["card_dark"]
    TABLE_LOW, TABLE_HOVER = t["table_low"], t["table_hover"]
    THEME_NAME = name

FONT_DISPLAY = ("Georgia", 25, "bold")
FONT_DISPLAY_SM = ("Georgia", 18, "bold")
FONT_BODY = ("Segoe UI", 13,)
FONT_BODY_B = ("Segoe UI", 13, "bold")
FONT_MONO = ("Consolas", 11,)
FONT_MONO_SM = ("Consolas", 11)
# Dashboard-specific fonts — easier to read than the small mono font used elsewhere.
FONT_DASH_LABEL = ("Segoe UI", 11, "normal")
FONT_DASH_TEXT  = ("Segoe UI", 12, "bold")
FONT_DASH_VALUE = ("Georgia", 22, "bold")

def CATEGORIES_STOCK(): return get_types("stock_categorie")
def MARQUES_STOCK(): return get_types("stock_marque")
def CATEGORIES_DEPENSE(): return get_types("depense_categorie")
def TYPES_LENTILLE(): return get_types("lentille_type")
def DUREES_LENTILLE(): return get_types("lentille_duree")


SERVICES_RDV = ["Consultation de contrôle", "Essayage monture", "Adaptation lentilles", "Réparation", "Autre"]
STATUTS_RDV = ["Prévu", "Confirmé", "Terminé", "Annulé"]
MODES_PAIEMENT = ["Espèces", "TPE"]
STATUTS_COMMANDE = ["En attente", "Prête", "Livrée", "Annulée"]
JOURS_SEMAINE = ["L", "ME", "MA", "J", "V", "S"]
TYPE_CATEGORIES = [
    ("stock_categorie", "Catégories d'articles (Stock)"),
    ("stock_marque", "Marques (Stock)"),
    ("depense_categorie", "Catégories de dépenses"),
    ("lentille_type", "Types de lentilles"),
    ("lentille_duree", "Durées de lentilles"),
]


def today_iso():
    return datetime.date.today().isoformat()

def expense_is_paid_this_month(x):
    """For a recurring charge, 'paid' means paid *for the current month*.
    For a one-off charge, 'paid' just means paid. Rows without paye_mois
    (older data) are treated as not-yet-paid-for-this-month if recurring."""
    try:
        paid = bool(x["paye"])
    except (KeyError, IndexError, TypeError):
        paid = False
    if not paid:
        return False
    try:
        recurring = bool(x["recurrent"])
    except (KeyError, IndexError, TypeError):
        recurring = False
    if not recurring:
        return True
    try:
        paye_mois = x["paye_mois"] or ""
    except (KeyError, IndexError, TypeError):
        paye_mois = ""
    return paye_mois == datetime.date.today().strftime("%Y-%m")

def money(n):
    try:
        n = float(n or 0)
    except (TypeError, ValueError):
        n = 0
    return f"{n:,.0f}".replace(",", " ") + " DH"


UNITES = ["", "un", "deux", "trois", "quatre", "cinq", "six", "sept", "huit", "neuf", "dix",
          "onze", "douze", "treize", "quatorze", "quinze", "seize"]
DIZAINES = ["", "dix", "vingt", "trente", "quarante", "cinquante", "soixante", "soixante-dix",
            "quatre-vingt", "quatre-vingt-dix"]


def _under_1000(num):
    out = ""
    c, r = divmod(num, 100)
    if c > 0:
        out += (UNITES[c] + " cent" if c > 1 else "cent")
        if r == 0 and c > 1:
            out += "s"
        if r > 0:
            out += " "
    if r > 0:
        if r < 17:
            out += UNITES[r]
        elif r < 20:
            out += "dix-" + UNITES[r - 10]
        else:
            d, u = divmod(r, 10)
            if d in (7, 9):
                out += DIZAINES[d - 1] + "-" + ("et-onze" if u == 1 else UNITES[10 + u])
            else:
                if u > 0:
                    out += DIZAINES[d] + ("-et-" + UNITES[u] if u == 1 and d != 8 else "-" + UNITES[u])
                else:
                    out += DIZAINES[d] + ("s" if d == 8 else "")
    return out


def num_to_french_words(n):
    n = int(round(float(n or 0)))
    if n == 0:
        return "zéro"
    parts = []
    millions, rest = divmod(n, 1_000_000)
    milliers, reste = divmod(rest, 1000)
    if millions > 0:
        parts.append(_under_1000(millions) + (" millions" if millions > 1 else " million"))
    if milliers > 0:
        parts.append("mille" if milliers == 1 else _under_1000(milliers) + " mille")
    if reste > 0 or not parts:
        parts.append(_under_1000(reste))
    return " ".join(p for p in parts if p).strip()


def _ics_escape(text):
    return (text or "").replace("\\", "\\\\").replace(",", "\\,").replace(";", "\\;").replace("\n", "\\n")


def generate_ics(events):
    """Construit un fichier calendrier standard (.ics). Compatible avec Google Agenda,
    le Calendrier Samsung, Outlook, etc. — n'importe quelle app calendrier sait
    l'importer, sur ordinateur comme sur téléphone."""
    lines = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//OPTILUX//FR", "CALSCALE:GREGORIAN"]
    stamp = datetime.datetime.utcnow().strftime("%Y%m%dT%H%M%SZ")
    for e in events:
        lines.append("BEGIN:VEVENT")
        lines.append(f"UID:{e['uid']}@optilux.local")
        lines.append(f"DTSTAMP:{stamp}")
        lines.append(f"DTSTART:{e['start'].strftime('%Y%m%dT%H%M%S')}")
        lines.append(f"DTEND:{e['end'].strftime('%Y%m%dT%H%M%S')}")
        lines.append(f"SUMMARY:{_ics_escape(e['summary'])}")
        if e.get("description"):
            lines.append(f"DESCRIPTION:{_ics_escape(e['description'])}")
        lines.append("BEGIN:VALARM")
        lines.append("TRIGGER:-PT30M")
        lines.append("ACTION:DISPLAY")
        lines.append("DESCRIPTION:Rappel")
        lines.append("END:VALARM")
        lines.append("END:VEVENT")
    lines.append("END:VCALENDAR")
    return "\r\n".join(lines) + "\r\n"


def open_file(path):
    """Ouvre un fichier avec l'application par défaut du système."""
    try:
        if sys.platform.startswith("win"):
            os.startfile(path)  # noqa
        elif sys.platform == "darwin":
            subprocess.run(["open", path])
        else:
            subprocess.run(["xdg-open", path])
    except Exception:
        messagebox.showinfo("Fichier généré", f"Le fichier a été enregistré ici :\n{path}")


# ---------------------------------------------------------------- STYLE
def setup_style(root):
    style = ttk.Style(root)
    try:
        style.theme_use("clam")
    except tk.TclError:
        pass
    style.configure("Treeview", background=SURFACE, fieldbackground=SURFACE, foreground=TEXT,
                     rowheight=38, font=FONT_BODY, borderwidth=0)
    style.configure("Treeview.Heading", background=CONTENT_BG, foreground=SLATE,
                     font=FONT_MONO_SM, borderwidth=0, relief="flat")
    style.map("Treeview.Heading", background=[("active", CONTENT_BG)])
    style.map("Treeview", background=[("selected", TABLE_SELECTED)], foreground=[("selected", TEXT)])
    style.configure("TNotebook", background=CONTENT_BG, borderwidth=0)

    style.configure("Themed.TCombobox", fieldbackground=SURFACE, background=SURFACE,
                     foreground=TEXT, arrowcolor=TEXT, bordercolor=BORDER, borderwidth=1)
    style.map("Themed.TCombobox", fieldbackground=[("readonly", SURFACE)],
              foreground=[("readonly", TEXT)])
    root.option_add("*TCombobox*Listbox.background", SURFACE)
    root.option_add("*TCombobox*Listbox.foreground", TEXT)
    root.option_add("*TCombobox*Listbox.selectBackground", RED)
    root.option_add("*TCombobox*Listbox.selectForeground", WHITE)

    # Branded scrollbar (red thumb / paper trough) — used everywhere via style="Custom.Vertical.TScrollbar"
    style.configure("Custom.Vertical.TScrollbar", background=RED, troughcolor=CONTENT_BG,
                     bordercolor=CONTENT_BG, arrowcolor=WHITE, relief="flat",
                     gripcount=0, arrowsize=12, width=12)
    style.map("Custom.Vertical.TScrollbar",
              background=[("active", RED_BRIGHT), ("pressed", RED)],
              arrowcolor=[("disabled", CONTENT_BG)])
    return style


# ---------------------------------------------------------------- COLOR / HOVER ANIMATION
def _hex_to_rgb(widget, value):
    """Resolve a Tk color string to (r, g, b). Handles '#rrggbb', '#rgb', and named
    colors like 'white', 'red', 'SystemButtonFace' via Tk's own resolver."""
    value = str(value)
    if value.startswith("#") and len(value) in (4, 7):
        h = value.lstrip("#")
        if len(h) == 3:
            h = "".join(c * 2 for c in h)
        return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))
    try:
        r, g, b = widget.winfo_rgb(value)  # each 0..65535
        return (r // 256, g // 256, b // 256)
    except tk.TclError:
        return (0, 0, 0)


def _rgb_to_hex(rgb):
    return "#%02x%02x%02x" % tuple(max(0, min(255, int(round(c)))) for c in rgb)


def _lerp_color(widget, c1, c2, t):
    r1, g1, b1 = _hex_to_rgb(widget, c1)
    r2, g2, b2 = _hex_to_rgb(widget, c2)
    return _rgb_to_hex((r1 + (r2 - r1) * t, g1 + (g2 - g1) * t, b1 + (b2 - b1) * t))


def _animate_props(widget, props, steps=8, duration=140):
    """props: list of (option_name, from_hex, to_hex). Smoothly animates widget.configure(...)."""
    gen = getattr(widget, "_anim_gen", 0) + 1
    widget._anim_gen = gen

    def step(i=0):
        if not widget.winfo_exists() or getattr(widget, "_anim_gen", None) != gen:
            return
        t = i / steps
        try:
             widget.configure(**{name: _lerp_color(widget, c1, c2, t) for name, c1, c2 in props})
        except tk.TclError:
            return
        if i < steps:
            widget.after(max(1, duration // steps), lambda: step(i + 1))
    step()


def fade_in_page(page, steps=6, duration=130):
    """Fades a page's own text in — each label/button starts invisible (text color = its
    own background) and animates to its real color. Scoped entirely to this page's widgets;
    never touches window transparency, so the desktop is never visible behind the app."""
    targets = []

    def collect(w):
        try:
            keys = w.keys()
        except Exception:
            keys = ()
        if "fg" in keys:
            try:
                bg = str(w.cget("bg"))
                real_fg = str(w.cget("fg"))
                if real_fg != bg:
                    w.configure(fg=bg)
                    targets.append((w, bg, real_fg))
            except Exception:
                pass
        for c in w.winfo_children():
            collect(c)

    collect(page)
    if not targets:
        return

    gen = getattr(page, "_content_fade_gen", 0) + 1
    page._content_fade_gen = gen

    def step(i=0):
        if not page.winfo_exists() or getattr(page, "_content_fade_gen", None) != gen:
            return
        t = (i + 1) / steps
        for w, bg, real_fg in targets:
            if not w.winfo_exists():
                continue
            try:
                w.configure(fg=_lerp_color(w, bg, real_fg, t))
            except Exception:
                pass
        if i + 1 < steps:
            page.after(max(1, duration // steps), lambda: step(i + 1))

    step(0)


def add_hover_animation(widget, base_bg, hover_bg, base_fg=None, hover_fg=None):
    """Smooth color-fade hover effect, since Tk's built-in 'active' state swap is instant."""
    if base_fg is None:
        base_fg = str(widget.cget("fg"))
    if hover_fg is None:
        hover_fg = base_fg
    widget.configure(activebackground=hover_bg, activeforeground=hover_fg)

    def on_enter(_e):
        props = [("bg", base_bg, hover_bg)]
        if hover_fg != base_fg:
            props.append(("fg", base_fg, hover_fg))
        _animate_props(widget, props)

    def on_leave(_e):
        props = [("bg", hover_bg, base_bg)]
        if hover_fg != base_fg:
            props.append(("fg", hover_fg, base_fg))
        _animate_props(widget, props)

    widget.bind("<Enter>", on_enter)
    widget.bind("<Leave>", on_leave)
    return widget


# ---------------------------------------------------------------- ICONS
# Small line-icons drawn programmatically (not emoji/fonts) so they render identically
# and reliably on any OS/Windows font set, in any brand color, at any size.
_ICON_CACHE = {}


def _ic_plus(d, s, c, lw):
    m = s * 0.2
    d.line([(s * 0.5, m), (s * 0.5, s - m)], fill=c, width=lw)
    d.line([(m, s * 0.5), (s - m, s * 0.5)], fill=c, width=lw)


def _ic_edit(d, s, c, lw):
    x1, y1, x2, y2 = s * 0.22, s * 0.80, s * 0.70, s * 0.30
    d.line([(x1, y1), (x2, y2)], fill=c, width=lw)
    d.line([(x2, y2), (x2 + s * 0.11, y2 - s * 0.02)], fill=c, width=lw)
    d.line([(x2, y2), (x2 - s * 0.02, y2 + s * 0.11)], fill=c, width=lw)
    r = lw * 0.65
    d.ellipse([x1 - r, y1 - r, x1 + r, y1 + r], fill=c)


def _ic_delete(d, s, c, lw):
    top = s * 0.30
    d.rectangle([s * 0.28, top, s * 0.72, s * 0.82], outline=c, width=lw)
    d.line([(s * 0.18, top), (s * 0.82, top)], fill=c, width=lw)
    d.line([(s * 0.40, top), (s * 0.40, top - s * 0.10)], fill=c, width=lw)
    d.line([(s * 0.60, top), (s * 0.60, top - s * 0.10)], fill=c, width=lw)
    d.line([(s * 0.40, top - s * 0.10), (s * 0.60, top - s * 0.10)], fill=c, width=lw)
    for fx in (0.44, 0.50, 0.56):
        d.line([(s * fx, top + s * 0.10), (s * fx, s * 0.74)], fill=c, width=max(1, lw - 1))


def _ic_document(d, s, c, lw):
    m, fold = s * 0.22, s * 0.16
    x0, y0 = m, s * 0.14
    x1, y1 = s - m - fold, s * 0.14
    x2, y2 = s - m, s * 0.14 + fold
    x3, y3 = s - m, s * 0.86
    x4, y4 = m, s * 0.86
    pts = [(x0, y0), (x1, y1), (x2, y2), (x3, y3), (x4, y4), (x0, y0)]
    for i in range(len(pts) - 1):
        d.line([pts[i], pts[i + 1]], fill=c, width=lw)
    d.line([(x1, y1), (x1, y2)], fill=c, width=lw)
    d.line([(x1, y2), (x2, y2)], fill=c, width=lw)
    for fy in (0.44, 0.58, 0.72):
        d.line([(m + s * 0.08, s * fy), (s - m - s * 0.08, s * fy)], fill=c, width=max(1, lw - 1))


def _ic_logout(d, s, c, lw):
    d.line([(s * 0.62, s * 0.18), (s * 0.30, s * 0.18)], fill=c, width=lw)
    d.line([(s * 0.30, s * 0.18), (s * 0.30, s * 0.82)], fill=c, width=lw)
    d.line([(s * 0.30, s * 0.82), (s * 0.62, s * 0.82)], fill=c, width=lw)
    d.line([(s * 0.42, s * 0.5), (s * 0.84, s * 0.5)], fill=c, width=lw)
    d.line([(s * 0.84, s * 0.5), (s * 0.68, s * 0.36)], fill=c, width=lw)
    d.line([(s * 0.84, s * 0.5), (s * 0.68, s * 0.64)], fill=c, width=lw)


def _ic_search(d, s, c, lw):
    r = s * 0.26
    cx, cy = s * 0.40, s * 0.40
    d.ellipse([cx - r, cy - r, cx + r, cy + r], outline=c, width=lw)
    d.line([(cx + r * 0.72, cy + r * 0.72), (s * 0.84, s * 0.84)], fill=c, width=lw)


def _ic_dashboard(d, s, c, lw):
    pad, gap = s * 0.16, s * 0.12
    cell = (s - 2 * pad - gap) / 2
    for x0, y0 in [(pad, pad), (pad + cell + gap, pad), (pad, pad + cell + gap), (pad + cell + gap, pad + cell + gap)]:
        d.rounded_rectangle([x0, y0, x0 + cell, y0 + cell], radius=cell * 0.22, outline=c, width=lw)


def _ic_clients(d, s, c, lw):
    cx, r = s * 0.5, s * 0.15
    d.ellipse([cx - r, s * 0.16, cx + r, s * 0.16 + 2 * r], outline=c, width=lw)
    d.arc([s * 0.18, s * 0.52, s * 0.82, s * 1.10], start=200, end=340, fill=c, width=lw)


def _ic_calendar(d, s, c, lw):
    m = s * 0.14
    d.rounded_rectangle([m, s * 0.22, s - m, s - m], radius=s * 0.08, outline=c, width=lw)
    d.line([(s * 0.30, s * 0.08), (s * 0.30, s * 0.28)], fill=c, width=lw)
    d.line([(s * 0.70, s * 0.08), (s * 0.70, s * 0.28)], fill=c, width=lw)
    d.line([(m, s * 0.42), (s - m, s * 0.42)], fill=c, width=lw)


def _ic_box(d, s, c, lw):
    m = s * 0.16
    d.rectangle([m, s * 0.32, s - m, s - m], outline=c, width=lw)
    d.line([(m, s * 0.32), (s * 0.5, s * 0.14)], fill=c, width=lw)
    d.line([(s - m, s * 0.32), (s * 0.5, s * 0.14)], fill=c, width=lw)
    d.line([(s * 0.5, s * 0.32), (s * 0.5, s - m)], fill=c, width=lw)


def _ic_sales(d, s, c, lw):
    m = s * 0.14
    d.rounded_rectangle([m, s * 0.28, s - m, s * 0.72], radius=s * 0.06, outline=c, width=lw)
    r = s * 0.12
    cx, cy = s * 0.5, s * 0.5
    d.ellipse([cx - r, cy - r, cx + r, cy + r], outline=c, width=max(1, lw - 1))


def _ic_wallet(d, s, c, lw):
    m = s * 0.16
    d.rounded_rectangle([m, s * 0.26, s - m, s * 0.78], radius=s * 0.07, outline=c, width=lw)
    d.rounded_rectangle([s * 0.60, s * 0.44, s * 0.82, s * 0.60], radius=s * 0.03, outline=c, width=max(1, lw - 1))


def _ic_users(d, s, c, lw):
    r = s * 0.13
    d.ellipse([s * 0.24 - r, s * 0.22, s * 0.24 + r, s * 0.22 + 2 * r], outline=c, width=lw)
    d.arc([s * 0.06, s * 0.56, s * 0.56, s * 1.05], start=200, end=340, fill=c, width=lw)
    r2 = s * 0.13
    d.ellipse([s * 0.64 - r2, s * 0.18, s * 0.64 + r2, s * 0.18 + 2 * r2], outline=c, width=lw)
    d.arc([s * 0.44, s * 0.52, s * 0.96, s * 0.98], start=200, end=340, fill=c, width=lw)


def _ic_warning(d, s, c, lw):
    p1, p2, p3 = (s * 0.5, s * 0.14), (s * 0.88, s * 0.82), (s * 0.12, s * 0.82)
    d.line([p1, p2], fill=c, width=lw)
    d.line([p2, p3], fill=c, width=lw)
    d.line([p3, p1], fill=c, width=lw)
    d.line([(s * 0.5, s * 0.38), (s * 0.5, s * 0.62)], fill=c, width=lw)
    r = lw * 0.7
    d.ellipse([s * 0.5 - r, s * 0.68 - r, s * 0.5 + r, s * 0.68 + r], fill=c)


def _ic_check(d, s, c, lw):
    d.line([(s * 0.18, s * 0.52), (s * 0.42, s * 0.76)], fill=c, width=lw)
    d.line([(s * 0.42, s * 0.76), (s * 0.84, s * 0.24)], fill=c, width=lw)


def _ic_close(d, s, c, lw):
    d.line([(s * 0.24, s * 0.24), (s * 0.76, s * 0.76)], fill=c, width=lw)
    d.line([(s * 0.76, s * 0.24), (s * 0.24, s * 0.76)], fill=c, width=lw)


def _ic_clock(d, s, c, lw):
    r = s * 0.34
    cx, cy = s * 0.5, s * 0.5
    d.ellipse([cx - r, cy - r, cx + r, cy + r], outline=c, width=lw)
    d.line([(cx, cy), (cx, cy - r * 0.55)], fill=c, width=max(1, lw - 1))
    d.line([(cx, cy), (cx + r * 0.42, cy + r * 0.14)], fill=c, width=max(1, lw - 1))


def _ic_growth(d, s, c, lw):
    d.line([(s * 0.16, s * 0.82), (s * 0.84, s * 0.82)], fill=c, width=lw)
    d.line([(s * 0.16, s * 0.82), (s * 0.38, s * 0.54)], fill=c, width=lw)
    d.line([(s * 0.38, s * 0.54), (s * 0.56, s * 0.68)], fill=c, width=lw)
    d.line([(s * 0.56, s * 0.68), (s * 0.86, s * 0.26)], fill=c, width=lw)
    d.line([(s * 0.86, s * 0.26), (s * 0.68, s * 0.26)], fill=c, width=lw)
    d.line([(s * 0.86, s * 0.26), (s * 0.86, s * 0.44)], fill=c, width=lw)


def _ic_lock(d, s, c, lw):
    m = s * 0.28
    d.rounded_rectangle([m, s * 0.46, s - m, s * 0.86], radius=s * 0.05, outline=c, width=lw)
    d.arc([s * 0.32, s * 0.14, s * 0.68, s * 0.60], start=180, end=360, fill=c, width=lw)
    r = lw * 0.75
    d.ellipse([s * 0.5 - r, s * 0.62 - r, s * 0.5 + r, s * 0.62 + r], fill=c)


def _ic_sun(d, s, c, lw):
    import math
    r = s * 0.22
    cx, cy = s * 0.5, s * 0.5
    d.ellipse([cx - r, cy - r, cx + r, cy + r], outline=c, width=lw)
    for ang in range(0, 360, 45):
        rad = math.radians(ang)
        x1, y1 = cx + (r + s * 0.10) * math.cos(rad), cy + (r + s * 0.10) * math.sin(rad)
        x2, y2 = cx + (r + s * 0.22) * math.cos(rad), cy + (r + s * 0.22) * math.sin(rad)
        d.line([(x1, y1), (x2, y2)], fill=c, width=lw)


def _ic_moon(d, s, c, lw):
    r = s * 0.32
    cx, cy = s * 0.5, s * 0.5
    d.ellipse([cx - r, cy - r, cx + r, cy + r], fill=c)
    r2 = s * 0.28
    cx2, cy2 = cx + s * 0.17, cy - s * 0.07
    d.ellipse([cx2 - r2, cy2 - r2, cx2 + r2, cy2 + r2], fill=(0, 0, 0, 0))


def _ic_bell(d, s, c, lw):
    cx = s * 0.5
    d.arc([s * 0.24, s * 0.16, s * 0.76, s * 0.78], start=180, end=360, fill=c, width=lw)
    d.line([(s * 0.24, s * 0.56), (s * 0.20, s * 0.74)], fill=c, width=lw)
    d.line([(s * 0.76, s * 0.56), (s * 0.80, s * 0.74)], fill=c, width=lw)
    d.line([(s * 0.20, s * 0.74), (s * 0.80, s * 0.74)], fill=c, width=lw)
    r = lw * 0.9
    d.ellipse([cx - r, s * 0.80 - r, cx + r, s * 0.80 + r], fill=c)


def _ic_package(d, s, c, lw):
    m = s * 0.16
    d.rectangle([m, s * 0.30, s - m, s - m], outline=c, width=lw)
    d.line([(m, s * 0.30), (s * 0.5, s * 0.14)], fill=c, width=lw)
    d.line([(s - m, s * 0.30), (s * 0.5, s * 0.14)], fill=c, width=lw)
    d.line([(s * 0.5, s * 0.30), (s * 0.5, s - m)], fill=c, width=lw)
    d.line([(m, s * 0.58), (s - m, s * 0.58)], fill=c, width=max(1, lw - 1))


def _ic_gear(d, s, c, lw):
    import math
    cx, cy, r = s * 0.5, s * 0.5, s * 0.20
    d.ellipse([cx - r, cy - r, cx + r, cy + r], outline=c, width=lw)
    r2 = lw * 0.7
    d.ellipse([cx - r2, cy - r2, cx + r2, cy + r2], fill=c)
    for ang in range(0, 360, 45):
        rad = math.radians(ang)
        x1, y1 = cx + (r + s * 0.04) * math.cos(rad), cy + (r + s * 0.04) * math.sin(rad)
        x2, y2 = cx + (r + s * 0.16) * math.cos(rad), cy + (r + s * 0.16) * math.sin(rad)
        d.line([(x1, y1), (x2, y2)], fill=c, width=lw)


_ICON_DRAWERS = {
    "plus": _ic_plus, "edit": _ic_edit, "delete": _ic_delete, "document": _ic_document,
    "logout": _ic_logout, "search": _ic_search, "dashboard": _ic_dashboard, "clients": _ic_clients,
    "calendar": _ic_calendar, "box": _ic_box, "sales": _ic_sales, "wallet": _ic_wallet,
    "users": _ic_users, "warning": _ic_warning, "check": _ic_check, "close": _ic_close,
    "clock": _ic_clock, "growth": _ic_growth, "lock": _ic_lock,
    "sun": _ic_sun, "moon": _ic_moon, "bell": _ic_bell,
    "package": _ic_package, "gear": _ic_gear,
}


def get_icon(name, color, size=16):
    """Renders (supersampled + downscaled for anti-aliasing) and caches a small icon.
    Returns None if Pillow isn't available or the name is unknown — callers must handle
    that gracefully (button still works, just without the icon)."""
    key = (name, color, size)
    if key in _ICON_CACHE:
        return _ICON_CACHE[key]
    draw_fn = _ICON_DRAWERS.get(name)
    if draw_fn is None:
        return None
    try:
        from PIL import Image, ImageDraw, ImageTk
        scale = 4
        s = size * scale
        img = Image.new("RGBA", (s, s), (0, 0, 0, 0))
        d = ImageDraw.Draw(img)
        lw = max(2, int(s * 0.075))
        draw_fn(d, s, color, lw)
        img = img.resize((size, size), Image.LANCZOS)
        photo = ImageTk.PhotoImage(img)
        _ICON_CACHE[key] = photo
        return photo
    except Exception:
        return None


# ---------------------------------------------------------------- WIDGET HELPERS
def _attach_icon(button, icon, color, size=16):
    if not icon:
        return
    photo = get_icon(icon, color, size)
    if photo:
        current = button.cget("text")
        button.configure(image=photo, compound="left", text=(" " + current if current else current))
        button._icon_ref = photo  # prevent garbage collection


def primary_button(parent, text, command, icon="plus"):
    b = tk.Button(parent, text=text, command=command, bg=RED, fg=WHITE,
                   font=FONT_BODY_B, relief="flat", padx=16, pady=8, cursor="hand2",
                   borderwidth=0)
    add_hover_animation(b, RED, RED_BRIGHT)
    _attach_icon(b, icon, WHITE)
    return b


def outline_button(parent, text, command, icon=None):
    b = tk.Button(parent, text=text, command=command, bg=SURFACE, fg=TEXT,
                   font=FONT_BODY, relief="solid", borderwidth=1, padx=16, pady=8,
                   cursor="hand2", highlightbackground=BORDER)
    add_hover_animation(b, SURFACE, CONTENT_BG)
    _attach_icon(b, icon, TEXT)
    return b


def ghost_button(parent, text, command, fg=SLATE, icon=None):
    b = tk.Button(parent, text=text, command=command, bg=INK, fg=fg,
                   font=FONT_BODY, relief="flat", borderwidth=0, cursor="hand2", anchor="w",
                   padx=0, pady=6)
    add_hover_animation(b, INK, INK, base_fg=fg, hover_fg=WHITE)
    _attach_icon(b, icon, WHITE)
    return b


def danger_button(parent, text, command, icon="delete"):
    """Même gabarit que outline_button, mais en rouge — pour les actions de suppression."""
    b = tk.Button(parent, text=text, command=command, bg=SURFACE, fg=RED,
                   font=FONT_BODY_B, relief="solid", borderwidth=1,
                   padx=16, pady=8, cursor="hand2", highlightbackground=RED)
    add_hover_animation(b, SURFACE, RED, base_fg=RED, hover_fg=WHITE)
    _attach_icon(b, icon, RED)
    return b


def field(parent, label_text, initial="", show=None, width=None):
    tk.Label(parent, text=label_text.upper(), font=FONT_MONO_SM, fg=SLATE, bg=SURFACE,
              anchor="w").pack(fill="x", pady=(8, 2))
    entry = tk.Entry(parent, font=FONT_BODY, relief="solid", borderwidth=1, show=show,
                      bg=SURFACE, fg=TEXT, insertbackground=TEXT,
                      highlightbackground=BORDER, highlightcolor=RED, highlightthickness=1)
    if width:
        entry.config(width=width)
    entry.insert(0, "" if initial is None else str(initial))
    entry.pack(fill="x", ipady=4)
    return entry

class DateField(tk.Frame):
    """Three spinboxes (J / M / A). Behaves like an Entry: .get() returns
    'YYYY-MM-DD', or '' if incomplete. The '×' clears all three — used for
    optional échéance / prévue fields."""
    def __init__(self, parent, initial=""):
        super().__init__(parent, bg=SURFACE, highlightbackground=BORDER, highlightcolor=RED,
                         highlightthickness=1)
        d = m = y = None
        if initial:
            try:
                dt = datetime.date.fromisoformat(str(initial).strip())
                d, m, y = dt.day, dt.month, dt.year
            except (ValueError, TypeError):
                pass
        self.day_var = tk.StringVar(value=f"{d:02d}" if d else "")
        self.month_var = tk.StringVar(value=f"{m:02d}" if m else "")
        self.year_var = tk.StringVar(value=str(y) if y else "")

        def make_spin(var, from_, to, width):
            return tk.Spinbox(self, from_=from_, to=to, width=width, font=FONT_MONO_SM,
                              bg=SURFACE, fg=TEXT, buttonbackground=CONTENT_BG,
                              relief="flat", justify="center", textvariable=var,
                              highlightthickness=0, borderwidth=0, wrap=True)

        make_spin(self.day_var, 1, 31, 3).pack(side="left", padx=(8, 2), pady=5)
        tk.Label(self, text="/", font=FONT_MONO_SM, bg=SURFACE, fg=SLATE).pack(side="left")
        make_spin(self.month_var, 1, 12, 3).pack(side="left", padx=2)
        tk.Label(self, text="/", font=FONT_MONO_SM, bg=SURFACE, fg=SLATE).pack(side="left")
        make_spin(self.year_var, 1900, 2100, 5).pack(side="left", padx=(2, 8))

        clear = tk.Label(self, text="×", font=FONT_BODY_B, bg=SURFACE, fg=SLATE, cursor="hand2")
        clear.pack(side="right", padx=(0, 8))
        clear.bind("<Button-1>", lambda e: self.clear())

    def clear(self):
        self.day_var.set("")
        self.month_var.set("")
        self.year_var.set("")

    def get(self):
        d = self.day_var.get().strip()
        m = self.month_var.get().strip()
        y = self.year_var.get().strip()
        if not (d and m and y):
            return ""
        try:
            return datetime.date(int(y), int(m), int(d)).isoformat()
        except (ValueError, tk.TclError):
            return ""

    def set(self, value):
        try:
            dt = datetime.date.fromisoformat(str(value).strip())
            self.day_var.set(f"{dt.day:02d}")
            self.month_var.set(f"{dt.month:02d}")
            self.year_var.set(str(dt.year))
        except (ValueError, TypeError):
            self.clear()


def field_date(parent, label_text, initial=""):
    """Same shape as field(), but renders a DateField. .get() returns the ISO date."""
    tk.Label(parent, text=label_text.upper(), font=FONT_MONO_SM, fg=SLATE, bg=SURFACE,
              anchor="w").pack(fill="x", pady=(8, 2))
    df = DateField(parent, initial=initial)
    df.pack(fill="x")
    return df

def dropdown(parent, label_text, options, initial=None):
    tk.Label(parent, text=label_text.upper(), font=FONT_MONO_SM, fg=SLATE, bg=SURFACE,
              anchor="w").pack(fill="x", pady=(8, 2))
    var = tk.StringVar(value=initial if initial in options else (options[0] if options else ""))
    om = ttk.Combobox(parent, textvariable=var, values=options, state="readonly", font=FONT_BODY,
                       style="Themed.TCombobox")
    om.pack(fill="x", ipady=2)
    return var


def oui_non(parent, label_text, initial_oui=False):
    """Paire de cases OUI / NON (comme sur la fiche papier) — une seule des deux est cochée."""
    row = tk.Frame(parent, bg=SURFACE)
    row.pack(fill="x", pady=(8, 2))
    tk.Label(row, text=label_text.upper(), font=FONT_MONO_SM, fg=SLATE, bg=SURFACE).pack(side="left")
    var = tk.StringVar(value="oui" if initial_oui else "non")
    opts = tk.Frame(row, bg=SURFACE)
    opts.pack(side="right")
    for val, txt in (("oui", "OUI"), ("non", "NON")):
        tk.Radiobutton(opts, text=txt, variable=var, value=val, font=FONT_MONO_SM,
                        bg=SURFACE, fg=TEXT, selectcolor=SURFACE, activebackground=SURFACE,
                        activeforeground=RED).pack(side="left", padx=4)
    return var


def mini_entry(parent, row, col, initial=""):
    e = tk.Entry(parent, font=FONT_MONO_SM, width=7, relief="solid", borderwidth=1,
                 bg=SURFACE, fg=TEXT, insertbackground=TEXT, justify="center")
    e.insert(0, initial or "")
    e.grid(row=row, column=col, padx=2, pady=2, sticky="ew")
    return e


def enable_row_hover(tree, hover_bg=None):
    """Highlights the row under the cursor. Tk gives priority to the FIRST tag in an item's
    tag list, so we prepend/remove 'hover' rather than append it."""
    if hover_bg is None:
        hover_bg = TABLE_HOVER
    tree.tag_configure("hover", background=hover_bg)
    last = {"row": None}

    def _clear(row):
        if row and tree.exists(row):
            tags = [t for t in tree.item(row, "tags") if t != "hover"]
            tree.item(row, tags=tuple(tags))

    def on_motion(event):
        row = tree.identify_row(event.y)
        if row == last["row"]:
            return
        _clear(last["row"])
        if row:
            tags = tuple(t for t in tree.item(row, "tags") if t != "hover")
            tree.item(row, tags=("hover",) + tags)
        last["row"] = row

    def on_leave(_event):
        _clear(last["row"])
        last["row"] = None

    tree.bind("<Motion>", on_motion)
    tree.bind("<Leave>", on_leave)


# ================================================================== APP
class OptiluxApp(tk.Tk):
    def __init__(self):
        # Must run before any window is shown: without this, Windows groups the app under
        # python.exe's own icon in the taskbar instead of using our custom one.
        if sys.platform.startswith("win"):
            try:
                import ctypes
                ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID("Optilux.Desktop.App.1")
            except Exception:
                pass

        super().__init__()
        apply_theme(get_setting("theme", "light"))

        self.title("OPTILUX — Gestion")
        self.geometry("1180x740")
        self.minsize(1180, 740)
        self._center_window()
        self.configure(bg=CONTENT_BG)
        setup_style(self)
        self._set_window_icon()

        self.current_user = None  # sqlite3.Row

        self.container = tk.Frame(self, bg=CONTENT_BG)
        self.container.pack(fill="both", expand=True)

        self.login_frame = LoginFrame(self.container, self)
        self.main_frame = None

        self.splash = SplashScreen(self.container, self, on_done=self._show_login)
        self.splash.place(relx=0, rely=0, relwidth=1, relheight=1)
        self.splash.tkraise()

    def toggle_theme(self):
        new_name = "dark" if THEME_NAME == "light" else "light"
        apply_theme(new_name)
        set_setting("theme", new_name)
        setup_style(self)  # refresh ttk styles (Treeview/scrollbar/combobox) in place
        self.configure(bg=CONTENT_BG)
        self.container.configure(bg=CONTENT_BG)
        self.login_frame.refresh_theme()
        # plain tk widgets bake their colors in at creation time, so the workspace is
        # rebuilt from scratch to pick up the new palette — cheap, since it's all local SQLite
        if self.main_frame is not None:
            current_tab = getattr(self.main_frame, "current_tab", "dashboard")
            self.main_frame.destroy()
            self.main_frame = MainFrame(self.container, self)
            self.main_frame.place(relx=0, rely=0, relwidth=1, relheight=1)
            self.main_frame.show_page(current_tab)
            self.main_frame.tkraise()

    def _show_login(self):
        self.splash.place_forget()
        self.splash.destroy()
        self.login_frame.place(relx=0, rely=0, relwidth=1, relheight=1)
        self.login_frame.tkraise()

    def _center_window(self):
        """Center the main window on the user's screen at startup."""
        self.update_idletasks()
        w = self.winfo_width()
        h = self.winfo_height()
        if w <= 1 or h <= 1:
            w, h = 1180, 740  # not yet realized — use the requested size
        sw = self.winfo_screenwidth()
        sh = self.winfo_screenheight()
        x = max(0, (sw - w) // 2)
        y = max(0, (sh - h) // 2)
        self.geometry(f"{w}x{h}+{x}+{y}")

    def _set_window_icon(self):
        # .ico natif pour la barre des tâches Windows
        if sys.platform.startswith("win") and os.path.exists(ICON_ICO):
            try:
                self.iconbitmap(ICON_ICO)
            except Exception:
                pass
        # PNG (via iconphoto) pour un rendu correct sur toutes les plateformes
        icon_photo = load_logo_photo(64)
        if icon_photo is not None:
            self._icon_photo_ref = icon_photo  # garder une référence, sinon Tk la jette
            try:
                self.iconphoto(True, icon_photo)
            except Exception:
                pass

    # ---------------- auth ----------------
    def do_login(self, user_row):
        self.current_user = user_row
        log_action(user_row["username"], "Connexion")
        if self.main_frame is not None:
            self.main_frame.destroy()
        self.main_frame = MainFrame(self.container, self)
        self.main_frame.place(relx=0, rely=0, relwidth=1, relheight=1)
        self.main_frame.tkraise()
        self._start_notifications()
        self.after(3000, self._silent_update_check)

    def do_logout(self):
        log_action(self.current_user["username"], "Déconnexion")
        self._stop_notifications()
        self.current_user = None
        self.main_frame.place_forget()
        self.login_frame.reset()
        self.login_frame.place(relx=0, rely=0, relwidth=1, relheight=1)
        self.login_frame.tkraise()

    def is_admin(self):
        return self.current_user is not None and self.current_user["role"] == "admin"

    # ---------------- desktop notifications ----------------
    NOTIFY_INTERVAL_MS = 60_000  # revérifie chaque minute

    def _start_notifications(self):
        self._notified_appointments = set()
        self._notified_expenses = set()
        self._notified_commandes = set()
        self._notify_job = None
        self._run_notification_checks(first_run=True)

    def _stop_notifications(self):
        if getattr(self, "_notify_job", None) is not None:
            try:
                self.after_cancel(self._notify_job)
            except Exception:
                pass
            self._notify_job = None

    def _run_notification_checks(self, first_run=False):
        if self.current_user is None:
            return
        try:
            self._check_appointment_reminders()
            self._check_expense_reminders()
            self._check_commande_reminders()
            if first_run:
                self._check_low_stock_summary()
        except Exception as e:
            print("Erreur vérification notifications:", e)
        self._notify_job = self.after(self.NOTIFY_INTERVAL_MS, self._run_notification_checks)

    @staticmethod
    def find_client_phone(client_nom):
        """Les rendez-vous stockent le nom du client en texte libre (pas de lien direct
        vers la fiche client) — on retrouve son téléphone par correspondance de nom."""
        if not client_nom:
            return None
        conn = get_connection()
        row = conn.execute(
            "SELECT tel FROM clients WHERE TRIM(nom || ' ' || COALESCE(prenom,'')) = ? OR nom = ?",
            (client_nom.strip(), client_nom.strip())
        ).fetchone()
        conn.close()
        return row["tel"] if row and row["tel"] else None

    def _check_appointment_reminders(self):
        now = datetime.datetime.now()
        conn = get_connection()
        rows = conn.execute(
            "SELECT * FROM appointments WHERE date=? AND statut NOT IN ('Terminé','Annulé')",
            (now.date().isoformat(),)
        ).fetchall()
        conn.close()
        for a in rows:
            try:
                hh, mm = map(int, a["heure"].split(":"))
                apt_dt = now.replace(hour=hh, minute=mm, second=0, microsecond=0)
            except Exception:
                continue
            delta_min = (apt_dt - now).total_seconds() / 60
            if 0 <= delta_min <= 15 and a["id"] not in self._notified_appointments:
                notify("Rendez-vous à venir",
                       f"{a['client_nom']} à {a['heure']} — {a['service']}",
                       icon_path=ICON_ICO)
                phone = self.find_client_phone(a["client_nom"])
                if phone:
                    msg = (f"Bonjour M./Mme {a['client_nom']}, un rappel de votre rendez-vous "
                           f"chez OPTILUX aujourd'hui à {a['heure']} ({a['service']}). À bientôt !")
                    send_whatsapp(phone, msg)
                self._notified_appointments.add(a["id"])

    def _check_expense_reminders(self):
        """Rappel de charge : prévient 7 jours avant l'échéance d'une dépense, puis le jour même."""
        today = datetime.date.today()
        conn = get_connection()
        rows = conn.execute(
            "SELECT * FROM expenses WHERE date_echeance IS NOT NULL AND date_echeance != ''"
        ).fetchall()
        conn.close()
        for x in rows:
            try:
                echeance = datetime.date.fromisoformat(x["date_echeance"])
            except ValueError:
                continue
            if expense_is_paid_this_month(x):
                continue
            days_left = (echeance - today).days
            if days_left <= 7 and x["id"] not in self._notified_expenses:
                label = "aujourd'hui" if days_left == 0 else f"dans {days_left} jour(s)"
                notify("Rappel de charge",
                       f"{x['libelle']} ({money(x['montant'])}) — échéance {label}",
                       icon_path=ICON_ICO)
                self._notified_expenses.add(x["id"])

    def _check_commande_reminders(self):
        today = datetime.date.today()
        conn = get_connection()
        rows = conn.execute("""SELECT commandes.*, clients.nom as client_nom, clients.prenom as client_prenom
                                FROM commandes LEFT JOIN clients ON clients.id = commandes.client_id
                                WHERE statut NOT IN ('Livrée','Annulée') AND date_prevue IS NOT NULL""").fetchall()
        conn.close()
        for cmd in rows:
            try:
                prevue = datetime.date.fromisoformat(cmd["date_prevue"])
            except ValueError:
                continue
            days_left = (prevue - today).days
            if days_left <= CommandesPage.ALERT_DAYS and cmd["id"] not in self._notified_commandes:
                if (cmd["type"] or "client") == "fournisseur":
                    who_label = cmd["fournisseur"] or "Fournisseur"
                else:
                    who_label = f"{cmd['client_nom'] or ''} {cmd['client_prenom'] or ''}".strip() or "Client"
                when = "en retard" if days_left < 0 else ("aujourd'hui" if days_left == 0 else f"dans {days_left} jour(s)")
                notify("Commande à suivre",
                       f"{who_label} — {cmd['description'] or 'commande'} ({when})",
                       icon_path=ICON_ICO)
                self._notified_commandes.add(cmd["id"])

    def _check_low_stock_summary(self):
        conn = get_connection()
        rows = conn.execute("SELECT * FROM stock WHERE qte<=seuil").fetchall()
        conn.close()
        if not rows:
            return
        if len(rows) == 1:
            msg = f"{rows[0]['nom']} — {rows[0]['qte']} restant(s)"
        else:
            names = ", ".join(r["nom"] for r in rows[:3])
            extra = f" (+{len(rows) - 3})" if len(rows) > 3 else ""
            msg = f"{names}{extra}"
        notify("Stock faible", msg, icon_path=ICON_ICO)


# ================================================================== SPLASH / LOADING
class SplashScreen(tk.Frame):
    """Boot splash: the glasses logo fades in/out, with skeleton shapes suggesting the
    general shape of the app (sidebar + header + cards) while things get ready."""
    DURATION_MS = 1700

    def __init__(self, parent, app, on_done):
        super().__init__(parent, bg=INK)
        self.app = app
        self.on_done = on_done
        self._alive = True

        center = tk.Frame(self, bg=INK)
        center.place(relx=0.5, rely=0.38, anchor="center")

        self.frames = build_pulse_frames(max_size=88, bg_hex=INK, steps=18)
        self.logo_label = tk.Label(center, bg=INK, borderwidth=0)
        if self.frames:
            self.logo_label.configure(image=self.frames[0])
        else:
            static = load_logo_photo(88)
            self._static_ref = static
            if static:
                self.logo_label.configure(image=static)
        self.logo_label.pack()

        brand = tk.Frame(center, bg=INK)
        brand.pack(pady=(16, 0))
        tk.Label(brand, text="OPTI", font=FONT_DISPLAY_SM, bg=INK, fg=WHITE).pack(side="left")
        tk.Label(brand, text="LUX", font=FONT_DISPLAY_SM, bg=INK, fg=RED_BRIGHT).pack(side="left")

        # skeleton silhouette of the app shape: a sidebar block + header/card blocks
        skeleton = tk.Frame(self, bg=INK)
        skeleton.place(relx=0.5, rely=0.72, anchor="center")

        sk_row = tk.Frame(skeleton, bg=INK)
        sk_row.pack()
        self.sk_sidebar = tk.Frame(sk_row, bg="#1c1c1e", width=46, height=130)
        self.sk_sidebar.pack(side="left", padx=(0, 14))
        self.sk_sidebar.pack_propagate(False)

        sk_main = tk.Frame(sk_row, bg=INK)
        sk_main.pack(side="left")
        self.sk_bars = []
        for w, h in [(220, 16), (150, 10)]:
            b = tk.Frame(sk_main, bg="#1c1c1e", width=w, height=h)
            b.pack(anchor="w", pady=(0, 10))
            b.pack_propagate(False)
            self.sk_bars.append(b)
        cards_row = tk.Frame(sk_main, bg=INK)
        cards_row.pack(anchor="w")
        for _ in range(3):
            c = tk.Frame(cards_row, bg="#1c1c1e", width=64, height=46)
            c.pack(side="left", padx=(0, 10))
            c.pack_propagate(False)
            self.sk_bars.append(c)

        self._frame_idx = 0
        self._direction = 1
        self._pulse_logo()
        self._shimmer_phase = 0.0
        self._shimmer_skeleton()

        self.after(self.DURATION_MS, self._finish)

    def _pulse_logo(self):
        if not self._alive or not self.winfo_exists():
            return
        if self.frames:
            self.logo_label.configure(image=self.frames[self._frame_idx])
            self._frame_idx += self._direction
            last = len(self.frames) - 1
            if self._frame_idx >= last:
                self._frame_idx = last
                self._direction = -1
            elif self._frame_idx <= 0:
                self._frame_idx = 0
                self._direction = 1
        self.after(45, self._pulse_logo)

    def _shimmer_skeleton(self):
        if not self._alive or not self.winfo_exists():
            return
        import math
        self._shimmer_phase += 0.18
        for i, bar in enumerate([self.sk_sidebar] + self.sk_bars):
            phase = (math.sin(self._shimmer_phase + i * 0.5) + 1) / 2
            bar.configure(bg=_lerp_color(bar, "#1c1c1e", "#333336", phase))
        self.after(70, self._shimmer_skeleton)

    def _finish(self):
        self._alive = False
        if self.on_done:
            self.on_done()


# ================================================================== LOGIN
class LoginFrame(tk.Frame):
    def __init__(self, parent, app: OptiluxApp):
        super().__init__(parent, bg=INK)
        self.app = app
        self._build()

    def _build(self):
        for w in self.winfo_children():
            w.destroy()

        card = tk.Frame(self, bg=SURFACE, padx=40, pady=40, highlightbackground=BORDER, highlightthickness=1)
        card.place(relx=0.5, rely=0.5, anchor="center", width=400)

        brand = tk.Frame(card, bg=SURFACE)
        brand.pack(pady=(0, 22))
        self.logo_photo = load_logo_photo(56)
        if self.logo_photo is not None:
            tk.Label(brand, image=self.logo_photo, bg=SURFACE).pack(side="left", padx=(0, 10))
        tk.Label(brand, text="OPTI", font=FONT_DISPLAY, bg=SURFACE, fg=TEXT).pack(side="left")
        tk.Label(brand, text="LUX", font=FONT_DISPLAY, bg=SURFACE, fg=RED).pack(side="left")

        tk.Label(card, text="Connexion", font=FONT_DISPLAY_SM, bg=SURFACE, fg=TEXT).pack()
        tk.Label(card, text="GESTION INTERNE", font=FONT_MONO_SM, bg=SURFACE, fg=SLATE).pack(pady=(0, 20))

        self.user_entry = field(card, "Identifiant")
        self.pass_entry = field(card, "Mot de passe", show="•")

        self.error_lbl = tk.Label(card, text="Identifiant ou mot de passe incorrect.",
                                    font=FONT_MONO_SM, fg=RED, bg=SURFACE)

        btn = primary_button(card, "Se connecter", self.attempt_login, icon="lock")
        btn.pack(fill="x", pady=(18, 6), ipady=4)

        tk.Label(card, text="Démo — Admin : admin / admin123", font=FONT_MONO_SM,
                  fg=SLATE, bg=SURFACE, justify="center").pack(pady=(14, 0))
        tk.Label(card, text="Employé : employe1 / employe123", font=FONT_MONO_SM,
                  fg=SLATE, bg=SURFACE, justify="center").pack()

        self.user_entry.bind("<Return>", lambda e: self.attempt_login())
        self.pass_entry.bind("<Return>", lambda e: self.attempt_login())

        self.user_entry.bind("<Key>", lambda e: self.error_lbl.pack_forget())
        self.pass_entry.bind("<Key>", lambda e: self.error_lbl.pack_forget())

        self.user_entry.focus_set()

    def reset(self):
        self.user_entry.delete(0, "end")
        self.pass_entry.delete(0, "end")
        self.error_lbl.pack_forget()
        self.user_entry.focus_set()

    def refresh_theme(self):
        """Rebuild with current theme colors — called after a theme toggle so the login
        screen stays correct even if the user never sees it again until after logout."""
        self._build()

    def reset(self):
        self.user_entry.delete(0, "end")
        self.pass_entry.delete(0, "end")
        self.error_lbl.pack_forget()

    def attempt_login(self):
        from auth import verify_password
        username = self.user_entry.get().strip()
        password = self.pass_entry.get()

        # If the user pressed Enter from the username field (or clicked the button)
        # with no password yet, nudge them to the password field instead of
        # showing the generic "wrong credentials" error.
        if not password:
            self.error_lbl.configure(text="Mot de passe requis.")
            self.error_lbl.pack(pady=(4, 0))
            self.pass_entry.focus_set()
            return

        conn = get_connection()
        row = conn.execute("SELECT * FROM users WHERE username=?", (username,)).fetchone()
        conn.close()
        if row and verify_password(password, row["password_hash"], row["salt"]):
            self.app.do_login(row)
        else:
            self.error_lbl.configure(text="Identifiant ou mot de passe incorrect.")
            self.error_lbl.pack(pady=(4, 0))


# ================================================================== MAIN FRAME (sidebar + pages)
class MainFrame(tk.Frame):
    def __init__(self, parent, app: OptiluxApp):
        super().__init__(parent, bg=CONTENT_BG)
        self.app = app

        # ---------------- sidebar ----------------
        sidebar = tk.Frame(self, bg=INK, width=248)
        sidebar.pack(side="left", fill="y")
        sidebar.pack_propagate(False)

        brand = tk.Frame(sidebar, bg=INK)
        brand.pack(pady=26)
        self.logo_photo = load_logo_photo(44)
        if self.logo_photo is not None:
            tk.Label(brand, image=self.logo_photo, bg=INK).pack(side="left", padx=(0, 8))
        tk.Label(brand, text="OPTI", font=FONT_DISPLAY_SM, bg=INK, fg=WHITE).pack(side="left")
        tk.Label(brand, text="LUX", font=FONT_DISPLAY_SM, bg=INK, fg=RED_BRIGHT).pack(side="left")

        self.nav_buttons = {}
        nav_items = [
            ("dashboard", "Tableau de bord", "dashboard", False),
            ("clients", "Clients", "clients", False),
            ("appointments", "Rendez-vous", "calendar", False),
            ("stock", "Stock", "box", False),
            ("sales", "Ventes", "sales", False),
            ("commandes", "Commandes", "package", False),
            ("factures", "Factures", "document", False),
            ("expenses", "Dépenses", "wallet", True),
            ("types", "Types & Produits", "gear", True),
            ("users", "Utilisateurs", "users", True),
        ]
        for key, label, icon_name, admin_only in nav_items:
            if admin_only and not app.is_admin():
                continue
            b = tk.Button(sidebar, text=" " + label.upper(), font=FONT_MONO_SM, bg=INK, fg=SLATE,
                          bd=0, anchor="w", padx=24, pady=10, cursor="hand2",
                          command=lambda k=key: self.show_page(k))
            icon = get_icon(icon_name, WHITE, size=15)
            if icon:
                b.configure(image=icon, compound="left")
                b._icon_ref = icon
            b._rest_bg = INK
            self._add_nav_hover(b, hover_bg=CHARCOAL, hover_fg=WHITE)
            b.pack(fill="x")
            self.nav_buttons[key] = b

        footer = tk.Frame(sidebar, bg=INK)
        footer.pack(side="bottom", fill="x", pady=20, padx=20)
        role_label = "Administrateur" if app.is_admin() else "Employé"
        tk.Label(footer, text=app.current_user["nom"], font=FONT_MONO_SM, bg=INK, fg=WHITE,
                  anchor="w").pack(fill="x")
        tk.Label(footer, text=role_label, font=FONT_MONO_SM, bg=INK,
                  fg=RED_BRIGHT if app.is_admin() else SLATE, anchor="w").pack(fill="x")
        theme_label = " Mode sombre" if THEME_NAME == "light" else " Mode clair"
        theme_icon = "moon" if THEME_NAME == "light" else "sun"
        ghost_button(footer, theme_label, app.toggle_theme, fg=SLATE, icon=theme_icon).pack(fill="x", pady=(14, 0))
        ghost_button(footer, " Vérifier les mises à jour", self.check_updates, fg=SLATE, icon="gear").pack(fill="x", pady=(4, 0))
        ghost_button(footer, " Déconnexion", app.do_logout, fg=SLATE, icon="logout").pack(fill="x", pady=(4, 0))

        # ---------------- content area ----------------
        content_wrap = tk.Frame(self, bg=CONTENT_BG)
        content_wrap.pack(side="left", fill="both", expand=True)
        self.content = tk.Frame(content_wrap, bg=CONTENT_BG, padx=30, pady=26)
        self.content.pack(fill="both", expand=True)

        self.pages = {}
        page_classes = {
            "dashboard": DashboardPage,
            "clients": ClientsPage,
            "appointments": AppointmentsPage,
            "stock": StockPage,
            "sales": SalesPage,
            "commandes": CommandesPage,
            "factures": FacturesPage,
        }
        if app.is_admin():
            page_classes["expenses"] = ExpensesPage
            page_classes["types"] = TypesPage
            page_classes["users"] = UsersPage

        for key, cls in page_classes.items():
            page = cls(self.content, app)
            self.pages[key] = page

        self.show_page("dashboard")
        self._declutter_gen = 0
        self._configure_bind_id = self.winfo_toplevel().bind("<Configure>", self._on_window_configure, add="+")

    def destroy(self):
        # sans ça, une ancienne instance de MainFrame (après déconnexion ou changement de
        # thème) continue de recevoir les événements <Configure> de la fenêtre principale
        # et plante en tentant d'agir sur des widgets déjà détruits.
        try:
            self.winfo_toplevel().unbind("<Configure>", self._configure_bind_id)
        except Exception:
            pass
        super().destroy()

    def _on_window_configure(self, event):
        """Pendant un redimensionnement actif de la fenêtre, on masque temporairement le(s)
        tableau(x) de la page active (le plus coûteux à recalculer, surtout avec beaucoup de
        colonnes/lignes) et on les restaure dès que le redimensionnement se stabilise. Le
        résultat final est strictement identique — seul le travail redondant pendant le
        glissement de la souris est évité."""
        if not self.winfo_exists():
            return
        try:
            top = self.winfo_toplevel()
        except Exception:
            return
        if event.widget is not top:
            return
        try:
            page = self.pages.get(self.current_tab)
        except Exception:
            return
        if not page:
            return
        trees = [w for w in (getattr(page, "tree", None), getattr(page, "audit_tree", None)) if w is not None]
        if not trees:
            return
        for t in trees:
            wrap = t.master  # le Frame bordé créé par make_table() (tableau + scrollbar)
            if wrap.winfo_ismapped() and not hasattr(wrap, "_declutter_info"):
                wrap._declutter_info = wrap.pack_info()
                wrap.pack_forget()
        self._declutter_gen += 1
        gen = self._declutter_gen
        self.after(150, lambda: self._restore_declutter(gen, trees))

    def _restore_declutter(self, gen, trees):
        if gen != self._declutter_gen:
            return  # un redimensionnement plus récent a pris le relais
        for t in trees:
            if not t.winfo_exists():
                continue
            wrap = t.master
            info = getattr(wrap, "_declutter_info", None)
            if info and wrap.winfo_exists():
                wrap.pack(**info)
                del wrap._declutter_info

    def show_page(self, key):
        if getattr(self, "current_tab", None) == key:
            return  # déjà sur cette page — pas de fondu inutile
        self.current_tab = key
        self._restyle_nav(key)  # instantané : le clic doit réagir tout de suite
        self._swap_content(key)

    def _restyle_nav(self, key):
        for k, b in self.nav_buttons.items():
            if not b.winfo_exists():
                continue
            is_active = (k == key)
            b._rest_bg = CHARCOAL if is_active else INK
            b._rest_fg = WHITE if is_active else SLATE
            b.configure(bg=b._rest_bg, fg=b._rest_fg)

    def _swap_content(self, key):
        page = self.pages.get(key)
        if page and page.winfo_exists():
            for k, p in self.pages.items():
                if k != key and p.winfo_exists() and p.winfo_ismapped():
                    p.place_forget()
            page.place(relx=0, rely=0, relwidth=1, relheight=1)
            page.tkraise()
            if hasattr(page, "refresh"):
                page.refresh()
            fade_in_page(page)

    @staticmethod
    def _add_nav_hover(widget, hover_bg, hover_fg):
        widget._rest_fg = str(widget.cget("fg"))

        def on_enter(_e):
            _animate_props(widget, [
                ("bg", str(widget.cget("bg")), hover_bg),
                ("fg", str(widget.cget("fg")), hover_fg),
            ])

        def on_leave(_e):
            rest_bg = getattr(widget, "_rest_bg", INK)
            rest_fg = getattr(widget, "_rest_fg", SLATE)
            _animate_props(widget, [
                ("bg", str(widget.cget("bg")), rest_bg),
                ("fg", str(widget.cget("fg")), rest_fg),
            ])

        widget.bind("<Enter>", on_enter)
        widget.bind("<Leave>", on_leave)

    def check_updates(self):
        from updater import check_for_update, download_and_apply

        try:
            info = check_for_update(UPDATE_URL, APP_VERSION)
        except Exception as e:
            messagebox.showerror("Mise à jour", f"Impossible de contacter le serveur :\n{e}")
            return

        if info is None:
            messagebox.showinfo("Mise à jour",
                                f"Vous utilisez déjà la version la plus récente ({APP_VERSION}).")
            return

        notes = info.get("notes", "")
        if not messagebox.askyesno(
            "Mise à jour disponible",
            f"Version {info['version']} disponible (vous avez {APP_VERSION}).\n\n"
            f"{notes}\n\nTélécharger et installer maintenant ?\n"
            "(L'application va se fermer — vos données ne sont pas touchées.)"
        ):
            return

        try:
            app_dir = os.path.dirname(os.path.abspath(__file__))
            n = download_and_apply(info, app_dir)
        except Exception as e:
            messagebox.showerror("Erreur de mise à jour", f"Échec du téléchargement :\n{e}")
            return

        messagebox.showinfo(
            "Mise à jour installée",
            f"{n} fichier(s) mis à jour vers la version {info['version']}.\n\n"
            "L'application va se fermer. Rouvrez-la pour utiliser la nouvelle version.")
        self.app.destroy()


# ================================================================== PAGE BASE
class BasePage(tk.Frame):
    def __init__(self, parent, app: OptiluxApp):
        super().__init__(parent, bg=CONTENT_BG)
        self.app = app

    def header(self, kicker, title, button_text=None, button_cmd=None):
        row = tk.Frame(self, bg=CONTENT_BG)
        row.pack(fill="x", pady=(0, 16))
        left = tk.Frame(row, bg=CONTENT_BG)
        left.pack(side="left")
        tk.Label(left, text=kicker.upper(), font=FONT_MONO_SM, fg=RED, bg=CONTENT_BG).pack(anchor="w")
        tk.Label(left, text=title, font=FONT_DISPLAY, fg=TEXT, bg=CONTENT_BG).pack(anchor="w")
        if button_text:
            clean_text = button_text.lstrip("+").strip()
            primary_button(row, clean_text, button_cmd, icon="plus").pack(side="right", anchor="e")
        tk.Frame(self, bg=BORDER, height=1).pack(fill="x", pady=(0, 16))

    def make_search_box(self, placeholder="Rechercher…"):
        wrap = tk.Frame(self, bg=SURFACE, highlightbackground=BORDER, highlightthickness=1)
        wrap.pack(fill="x", pady=(0, 12))
        icon = get_icon("search", SLATE, 15)
        icon_lbl = tk.Label(wrap, bg=SURFACE)
        if icon:
            icon_lbl.configure(image=icon)
            icon_lbl._icon_ref = icon
        icon_lbl.pack(side="left", padx=(10, 0))
        var = tk.StringVar()
        var.trace_add("write", lambda *a: self.refresh())
        entry = tk.Entry(wrap, textvariable=var, font=FONT_BODY, relief="flat",
                          bg=SURFACE, fg=TEXT, insertbackground=TEXT)
        entry.pack(fill="x", ipady=6, padx=6, side="left", expand=True)
        return var

    @staticmethod
    def row_matches(query, *values):
        if not query:
            return True
        q = query.lower()
        return any(q in str(v).lower() for v in values if v is not None)

    def mode_button(self, parent, text, command):
        """A two-state toggle button (e.g. 'Charges' vs 'Personnelles'). Uses a dynamic
        rest-color hover (reads b._rest_bg live) instead of a hover baked in at creation
        time, since the active/inactive color is switched after the fact by the caller."""
        b = tk.Button(parent, text=text, command=command, font=FONT_BODY_B, relief="solid",
                      borderwidth=1, padx=16, pady=8, cursor="hand2")
        b._rest_bg, b._rest_fg = SURFACE, TEXT

        def on_enter(_e):
            target = RED_BRIGHT if b._rest_bg == RED else CONTENT_BG
            _animate_props(b, [("bg", str(b.cget("bg")), target)])

        def on_leave(_e):
            _animate_props(b, [("bg", str(b.cget("bg")), b._rest_bg)])

        b.bind("<Enter>", on_enter)
        b.bind("<Leave>", on_leave)
        return b

    @staticmethod
    def set_mode_button_active(button, active):
        button._rest_bg = RED if active else SURFACE
        button._rest_fg = WHITE if active else TEXT
        button.configure(bg=button._rest_bg, fg=button._rest_fg,
                          highlightbackground=RED if active else BORDER)

    def show_detail_modal(self, title, build_content_fn, on_edit, on_history=None, width=480):
        """Fenêtre de consultation en lecture seule : affiche les infos proprement,
        avec un bouton Modifier explicite pour passer au formulaire d'édition —
        on ne tombe jamais directement en mode modification."""
        modal = ModalForm(self.app, title, width=width)
        build_content_fn(modal.body)

        def go_edit():
            modal.destroy()
            on_edit()

        if on_history is not None:
            def go_history():
                modal.destroy()
                on_history()
            outline_button(modal.footer, "Historique", go_history, icon="clock").pack(
                side="left", ipady=4, padx=(0, 8))

        primary_button(modal.footer, "Modifier", go_edit, icon="edit").pack(
            side="left", fill="x", expand=True, ipady=4, padx=(0, 8))
        outline_button(modal.footer, "Fermer", modal.destroy, icon="close").pack(side="left", ipady=4)
        return modal

    @staticmethod
    def info_row(parent, label, value, value_color=None):
        row = tk.Frame(parent, bg=SURFACE)
        row.pack(fill="x", pady=4)
        tk.Label(row, text=label.upper(), font=FONT_MONO_SM, bg=SURFACE, fg=SLATE).pack(side="left")
        tk.Label(row, text=value if value not in (None, "") else "—", font=FONT_BODY_B, bg=SURFACE,
                  fg=value_color or TEXT, wraplength=260, justify="right").pack(side="right")

    @staticmethod
    def section_label(parent, text):
        tk.Label(parent, text=text, font=FONT_MONO_SM, fg=RED, bg=SURFACE).pack(anchor="w", pady=(14, 4))

    def make_table(self, columns, height=14):
        wrap = tk.Frame(self, bg=SURFACE, highlightbackground=BORDER, highlightthickness=1)
        wrap.pack(fill="both", expand=True)
        tree = ttk.Treeview(wrap, columns=columns, show="headings", height=height)
        for c in columns:
            tree.heading(c, text=c.upper())
            tree.column(c, width=120, anchor="w")
        vsb = ttk.Scrollbar(wrap, orient="vertical", command=tree.yview, style="Custom.Vertical.TScrollbar")
        tree.configure(yscroll=vsb.set)
        tree.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")
        tree.tag_configure("odd", background=TABLE_ODD)
        tree.tag_configure("even", background=SURFACE)
        tree.tag_configure("low", background=TABLE_LOW, foreground=RED)
        enable_row_hover(tree)
        return tree


# ================================================================== MODAL BASE
class ModalForm(tk.Toplevel):
    def __init__(self, app, title, width=420):
        super().__init__(app)
        self.app = app
        self.title(title)
        self.configure(bg=SURFACE)
        self.geometry(f"{width}x560")
        self.minsize(width, 320)
        self.resizable(False, True)
        self.transient(app)
        self.grab_set()
        try:
            if sys.platform.startswith("win") and os.path.exists(ICON_ICO):
                self.iconbitmap(ICON_ICO)
        except Exception:
            pass

        # scrollable area (fields) — a Canvas + inner Frame, with a fixed footer below for Save/Cancel
        canvas_wrap = tk.Frame(self, bg=SURFACE)
        canvas_wrap.pack(fill="both", expand=True)
        self._canvas = tk.Canvas(canvas_wrap, bg=SURFACE, highlightthickness=0)
        vsb = ttk.Scrollbar(canvas_wrap, orient="vertical", command=self._canvas.yview,
                             style="Custom.Vertical.TScrollbar")
        self._canvas.configure(yscrollcommand=vsb.set)
        self._canvas.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")

        self.body = tk.Frame(self._canvas, bg=SURFACE, padx=24, pady=20)
        self._body_window = self._canvas.create_window((0, 0), window=self.body, anchor="nw")

        def _on_body_configure(_e):
            self._canvas.configure(scrollregion=self._canvas.bbox("all"))
        self.body.bind("<Configure>", _on_body_configure)

        def _on_canvas_configure(e):
            self._canvas.itemconfig(self._body_window, width=e.width)
        self._canvas.bind("<Configure>", _on_canvas_configure)

        def _on_wheel(e):
            delta = -1 if (e.num == 4 or e.delta > 0) else 1
            self._canvas.yview_scroll(delta, "units")
        # bind wheel only while the pointer is over this modal's canvas, so it doesn't
        # hijack scrolling in the main window behind it
        self._canvas.bind("<Enter>", lambda e: self._bind_wheel())
        self._canvas.bind("<Leave>", lambda e: self._unbind_wheel())

        tk.Label(self.body, text=title, font=FONT_DISPLAY_SM, bg=SURFACE, fg=TEXT).pack(anchor="w", pady=(0, 10))

        self.footer = tk.Frame(self, bg=SURFACE, padx=24, pady=16, highlightbackground=BORDER, highlightthickness=1)
        self.footer.pack(fill="x", side="bottom")

        self.protocol("WM_DELETE_WINDOW", self._on_close)
        self.center()

    def center(self):
        """Center this modal over the app window. Call again after any .geometry(...)
        change so the modal re-centers with its new size."""
        self.update_idletasks()
        try:
            pw = self.app.winfo_width()
            ph = self.app.winfo_height()
            px = self.app.winfo_rootx()
            py = self.app.winfo_rooty()
        except Exception:
            return
        if pw <= 1:
            pw = self.app.winfo_screenwidth()
        if ph <= 1:
            ph = self.app.winfo_screenheight()
        w = self.winfo_width()
        h = self.winfo_height()
        if w <= 1:
            w = self.winfo_reqwidth()
        if h <= 1:
            h = self.winfo_reqheight()
        x = px + max(0, (pw - w) // 2)
        y = py + max(0, (ph - h) // 2)
        self.geometry(f"+{x}+{y}")

    def _bind_wheel(self):
        self._canvas.bind_all("<MouseWheel>", self._on_mousewheel)
        self._canvas.bind_all("<Button-4>", self._on_mousewheel)
        self._canvas.bind_all("<Button-5>", self._on_mousewheel)

    def _unbind_wheel(self):
        self._canvas.unbind_all("<MouseWheel>")
        self._canvas.unbind_all("<Button-4>")
        self._canvas.unbind_all("<Button-5>")

    def _on_mousewheel(self, e):
        delta = -1 if (getattr(e, "num", None) == 4 or getattr(e, "delta", 0) > 0) else 1
        self._canvas.yview_scroll(delta, "units")

    def _on_close(self):
        self._unbind_wheel()
        self.destroy()

    def destroy(self):
        self._unbind_wheel()
        super().destroy()

    def buttons(self, on_save):
        row = self.footer
        primary_button(row, "Enregistrer", on_save, icon="check").pack(side="left", fill="x", expand=True, ipady=4, padx=(0, 8))
        outline_button(row, "Annuler", self._on_close, icon="close").pack(side="left", ipady=4)


# ================================================================== DASHBOARD
class DashboardPage(BasePage):
    def __init__(self, parent, app):
        super().__init__(parent, app)
        # Header lives directly on the page and is built once — it stays fixed at the
        # top while everything below it scrolls.
        self.header("Aperçu", "Tableau de bord")

        # Scrollable container below the header.
        self._canvas = tk.Canvas(self, bg=CONTENT_BG, highlightthickness=0)
        vsb = ttk.Scrollbar(self, orient="vertical", command=self._canvas.yview,
            style="Custom.Vertical.TScrollbar")
        self._canvas.configure(yscrollcommand=vsb.set)
        self._canvas.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")

        self.scroll_body = tk.Frame(self._canvas, bg=CONTENT_BG)
        self._body_window = self._canvas.create_window((0, 0), window=self.scroll_body, anchor="nw")

        def on_body_config(_e):
            self._canvas.configure(scrollregion=self._canvas.bbox("all"))
        self.scroll_body.bind("<Configure>", on_body_config)

        def on_canvas_config(e):
            # keep the inner frame as wide as the canvas so fill="x" works inside it
            self._canvas.itemconfig(self._body_window, width=e.width)
        self._canvas.bind("<Configure>", on_canvas_config)

        # Wheel only while the pointer is over the dashboard — mirrors ModalForm's pattern,
        # so opening a modal doesn't hijack scrolling here or vice versa.
        self._canvas.bind("<Enter>", lambda e: self._bind_wheel())
        self._canvas.bind("<Leave>", lambda e: self._unbind_wheel())

    def _bind_wheel(self):
        self._canvas.bind_all("<MouseWheel>", self._on_wheel)
        self._canvas.bind_all("<Button-4>", self._on_wheel)
        self._canvas.bind_all("<Button-5>", self._on_wheel)

    def _unbind_wheel(self):
        self._canvas.unbind_all("<MouseWheel>")
        self._canvas.unbind_all("<Button-4>")
        self._canvas.unbind_all("<Button-5>")

    def _on_wheel(self, e):
        delta = -1 if (getattr(e, "num", None) == 4 or getattr(e, "delta", 0) > 0) else 1
        self._canvas.yview_scroll(delta, "units")

    def refresh(self):
        for w in self.scroll_body.winfo_children():
            w.destroy()
        
        export_row = tk.Frame(self.scroll_body, bg=CONTENT_BG)
        export_row.pack(fill="x", pady=(0, 16))
        outline_button(export_row, "Exporter calendrier (.ics)", self.export_calendar, icon="calendar").pack(side="right")

        conn = get_connection()
        ym = today_iso()[:7]
        month_sales = conn.execute("SELECT * FROM sales WHERE date LIKE ?", (ym + "%",)).fetchall()
        ca = sum(s["vente"] or 0 for s in month_sales)
        avances = sum(s["avance"] or 0 for s in month_sales)
        all_sales = conn.execute("SELECT * FROM sales").fetchall()
        credit = sum((s["vente"] or 0) - (s["avance"] or 0) for s in all_sales)
        month_exp = conn.execute(
            "SELECT * FROM expenses WHERE personnelle=0 AND (recurrent=1 OR date LIKE ?)",
            (ym + "%",)
        ).fetchall()
        charges = sum(x["montant"] or 0 for x in month_exp)
        dispo = avances - charges

        today = today_iso()
        today_appts = conn.execute("SELECT * FROM appointments WHERE date=? ORDER BY heure", (today,)).fetchall()
        low_stock = conn.execute("SELECT * FROM stock WHERE qte<=seuil").fetchall()

        commandes_rows = conn.execute("""SELECT commandes.*, clients.nom as client_nom, clients.prenom as client_prenom
                                          FROM commandes LEFT JOIN clients ON clients.id = commandes.client_id
                                          WHERE statut NOT IN ('Livrée','Annulée')""").fetchall()
        commande_alerts = [c for c in commandes_rows if CommandesPage.alert_state(c)]
        commande_alerts.sort(key=lambda c: c["date_prevue"] or "")

        expense_alerts = self._expense_alerts(conn)

        recent_factures = conn.execute("SELECT * FROM factures ORDER BY date DESC LIMIT 4").fetchall() \
            if self.app.is_admin() else []
        conn.close()

        if self.app.is_admin():
            kpi_row = tk.Frame(self.scroll_body, bg=CONTENT_BG)
            kpi_row.pack(fill="x", pady=(0, 20))
            self._kpi(kpi_row, "CA du mois", money(ca), icon="sales")
            self._kpi(kpi_row, "Avances encaissées", money(avances), icon="wallet")
            self._kpi(kpi_row, "Crédit (reste)", money(credit), accent=True, icon="clock")
            self._kpi(kpi_row, "Charges fixes (mois)", money(charges), icon="document")

        # ---- cartes du bas : 2 par rangée (plus de place pour le texte agrandi) ----
        card_specs = []
        if self.app.is_admin():
            card_specs.append(("Disponible net", "growth", lambda card: self._build_dispo(card, dispo)))
            card_specs.append(("Chiffre d'affaires", "sales", lambda card: self._build_ca(card, ca)))
            card_specs.append(("Rendez-vous aujourd'hui", "calendar", lambda card: self._build_list(
            card, today_appts, "Aucun rendez-vous aujourd'hui.",
                lambda a: (f"{a['heure']} — {a['client_nom']}", a["statut"], SLATE))))
            card_specs.append(("Alertes stock", "warning", lambda card: self._build_list(
            card, low_stock, "Tout le stock est suffisant.",
                lambda s: (s["nom"], f"{s['qte']} restant(s)", RED))))
            card_specs.append(("Commandes à suivre", "package", lambda card: self._build_list(
            card, commande_alerts[:6], "Aucune commande urgente.",
                 lambda c: (
                    ((c["fournisseur"] if (c["type"] or "client") == "fournisseur"
                        else f"{c['client_nom'] or ''} {c['client_prenom'] or ''}".strip()) or "—")
                        + (f" — {c['description']}" if c["description"] else ""),
                    "EN RETARD" if CommandesPage.alert_state(c) == "overdue" else "BIENTÔT", RED))))
            card_specs.append(("Dépenses à payer", "wallet", lambda card: self._build_list(
            card, expense_alerts[:6], "Aucune échéance proche.",
            lambda pair: self._expense_alert_row(pair))))
        if self.app.is_admin():
            card_specs.append(("Factures récentes", "document", lambda card: self._build_list(
                card, recent_factures, "Aucune facture enregistrée.",
                lambda f: (f["titre"] or "—", money(f["montant"]), SLATE))))

        grid_wrap = tk.Frame(self.scroll_body, bg=CONTENT_BG)
        grid_wrap.pack(fill="both", expand=True)
        grid_wrap.grid_columnconfigure(0, weight=1)
        grid_wrap.grid_columnconfigure(1, weight=1)
        for i, (title, icon, builder) in enumerate(card_specs):
            row, col = divmod(i, 2)
            card = tk.Frame(grid_wrap, bg=SURFACE, highlightbackground=BORDER, highlightthickness=1, padx=20, pady=18)
            card.grid(row=row, column=col, sticky="nsew",
                      padx=(0, 8) if col == 0 else (8, 0), pady=(0, 16))
            self._card_title(card, title, icon)
            builder(card)

    EXPENSE_ALERT_DAYS = 7

    def _expense_alerts(self, conn):
        """All charges not yet paid for the current month, sorted: overdue first,
        then soonest due date, then charges with no due date."""
        rows = conn.execute("SELECT * FROM expenses").fetchall()
        today = datetime.date.today()
        out = []
        for x in rows:
            if expense_is_paid_this_month(x):
                continue
            due = x["date_echeance"] or ""
            if due:
                try:
                    days_left = (datetime.date.fromisoformat(due) - today).days
                except ValueError:
                    days_left = None
            else:
                days_left = None
            out.append((days_left, x))
        out.sort(key=lambda pair: (pair[0] is None, pair[0] if pair[0] is not None else 0))
        return out

    @staticmethod
    def _expense_alert_row(pair):
        days_left, x = pair
        if days_left is None:
            when, color = "à payer", SLATE
        elif days_left < 0:
            when, color = f"EN RETARD ({-days_left}j)", RED
        elif days_left <= 7:
            # due within a week (or today) → red
            when, color = ("AUJOURD'HUI" if days_left == 0 else f"dans {days_left}j"), RED
        else:
            when, color = f"dans {days_left}j", SLATE
        return (f"{x['libelle']} — {money(x['montant'])}", when, color)

    def _build_dispo(self, card, dispo):
        tk.Label(card, text="Avances − Charges fixes", font=FONT_DASH_TEXT, bg=SURFACE, fg=SLATE).pack(anchor="w", pady=(0, 10))
        tk.Label(card, text=money(dispo), font=FONT_DASH_VALUE, bg=SURFACE, fg=RED).pack(anchor="w")

    def _build_ca(self, card, ca):
        tk.Label(card, text="Total des ventes (Crédit inclue)", font=FONT_DASH_TEXT, bg=SURFACE, fg=SLATE).pack(anchor="w", pady=(0, 10))
        tk.Label(card, text=money(ca), font=FONT_DASH_VALUE, bg=SURFACE, fg=RED).pack(anchor="w")

    def _build_list(self, card, rows, empty_text, row_fn):
        if not rows:
            tk.Label(card, text=empty_text, font=FONT_DASH_TEXT, bg=SURFACE, fg=SLATE).pack(anchor="w")
            return
        for item in rows:
            label_txt, status_txt, status_color = row_fn(item)
            r = tk.Frame(card, bg=SURFACE)
            r.pack(fill="x", pady=4)
            if len(label_txt) > 26:
                tk.Label(r, text=label_txt, font=FONT_DASH_TEXT, bg=SURFACE, fg=TEXT, anchor="w",
                          wraplength=420, justify="left").pack(fill="x")
                tk.Label(r, text=status_txt, font=FONT_DASH_LABEL, bg=SURFACE, fg=status_color, anchor="w").pack(fill="x")
            else:
                tk.Label(r, text=label_txt, font=FONT_DASH_TEXT, bg=SURFACE, fg=TEXT).pack(side="left")
                tk.Label(r, text=status_txt, font=FONT_DASH_LABEL, bg=SURFACE, fg=status_color).pack(side="right")
    def export_calendar(self):
        conn = get_connection()
        today = today_iso()
        appts = conn.execute(
            "SELECT * FROM appointments WHERE date>=? AND statut NOT IN ('Terminé','Annulé') ORDER BY date, heure",
            (today,)).fetchall()
        commandes_rows = conn.execute("""SELECT commandes.*, clients.nom as client_nom, clients.prenom as client_prenom
                                          FROM commandes LEFT JOIN clients ON clients.id = commandes.client_id
                                          WHERE date_prevue>=? AND statut NOT IN ('Livrée','Annulée')
                                          ORDER BY date_prevue""", (today,)).fetchall()
        conn.close()

        events = []
        for a in appts:
            try:
                hh, mm = map(int, a["heure"].split(":"))
                start = datetime.datetime.strptime(a["date"], "%Y-%m-%d").replace(hour=hh, minute=mm)
            except Exception:
                continue
            events.append({
                "uid": f"optilux-rdv-{a['id']}",
                "start": start, "end": start + datetime.timedelta(minutes=30),
                "summary": f"RDV — {a['client_nom']}",
                "description": a["service"] or "",
            })
        for cmd in commandes_rows:
            try:
                start = datetime.datetime.strptime(cmd["date_prevue"], "%Y-%m-%d").replace(hour=9, minute=0)
            except Exception:
                continue
            who = (cmd["fournisseur"] if (cmd["type"] or "client") == "fournisseur"
                   else f"{cmd['client_nom'] or ''} {cmd['client_prenom'] or ''}".strip()) or "—"
            events.append({
                "uid": f"optilux-cmd-{cmd['id']}",
                "start": start, "end": start + datetime.timedelta(hours=1),
                "summary": f"Commande — {who}",
                "description": cmd["description"] or "",
            })

        if not events:
            messagebox.showinfo("Calendrier", "Aucun rendez-vous ou commande à venir à exporter.")
            return

        default_name = f"optilux_calendrier_{today}.ics"
        path = filedialog.asksaveasfilename(
            title="Enregistrer le calendrier", defaultextension=".ics", initialfile=default_name,
            filetypes=[("Calendrier iCalendar", "*.ics")])
        if not path:
            return
        try:
            with open(path, "w", encoding="utf-8") as f:
                f.write(generate_ics(events))
        except Exception as e:
            messagebox.showerror("Erreur", f"Impossible d'enregistrer le fichier :\n{e}")
            return

        log_action(self.app.current_user["username"], f"Calendrier exporté ({len(events)} événement(s))")
        messagebox.showinfo(
            "Calendrier exporté",
            f"{len(events)} événement(s) exporté(s) avec succès.\n\n"
            "Pour les voir sur votre téléphone (Samsung) : envoyez-vous ce fichier par e-mail, ou "
            "enregistrez-le dans Google Drive / OneDrive, puis ouvrez-le depuis votre téléphone — "
            "il s'ajoutera automatiquement à Google Agenda ou à l'application Calendrier Samsung.\n\n"
            "Astuce : refaites cet export de temps en temps pour garder votre téléphone à jour "
            "(ce n'est pas une synchronisation automatique en continu)."
        )
        open_file(path)

    def _card_title(self, parent, text, icon_name):
        icon = get_icon(icon_name, RED, size=17)
        lbl = tk.Label(parent, text=" " + text, font=FONT_DISPLAY_SM, bg=SURFACE, fg=TEXT,
                        compound="left")
        if icon:
            lbl.configure(image=icon)
            lbl._icon_ref = icon
        lbl.pack(anchor="w", pady=(0, 10))

    def _kpi(self, parent, label, value, accent=False, icon=None):
        bg = RED if accent else CARD_DARK
        card = tk.Frame(parent, bg=bg, padx=18, pady=16)
        card.pack(side="left", fill="both", expand=True, padx=6)
        top = tk.Frame(card, bg=bg)
        top.pack(fill="x")
        tk.Label(top, text=label.upper(), font=FONT_DASH_LABEL, bg=bg,
                  fg="white" if accent else SLATE).pack(side="left")
        icon_photo = get_icon(icon, WHITE, size=16) if icon else None
        if icon_photo:
            il = tk.Label(top, image=icon_photo, bg=bg)
            il._icon_ref = icon_photo
            il.pack(side="right")
        tk.Label(card, text=value, font=FONT_DASH_VALUE, bg=bg, fg=WHITE).pack(anchor="w", pady=(6, 0))


# ================================================================== CLIENTS
# ================================================================== CLIENTS
class ClientsPage(BasePage):
    def __init__(self, parent, app):
        super().__init__(parent, app)
        self.header("Fiches", "Clients", "+ Nouveau client", self.open_form)
        self.search_var = self.make_search_box("Rechercher un client…")
        self.tree = self.make_table(["Nom", "Prénom", "Téléphone", "Ville", "Remise", "Dernière Rx"])
        self.tree.bind("<Double-1>", lambda e: self.open_detail(self._selected_id()))

        btns = tk.Frame(self, bg=CONTENT_BG)
        btns.pack(fill="x", pady=(8, 0))
        outline_button(btns, "Voir la fiche", lambda: self.open_detail(self._selected_id()), icon="search").pack(side="left")
        danger_button(btns, "Supprimer", self.delete_selected).pack(side="left", padx=(0, 12))

    def _selected_id(self):
        sel = self.tree.selection()
        return int(sel[0]) if sel else None

    def refresh(self):
        for i in self.tree.get_children():
            self.tree.delete(i)
        q = self.search_var.get()
        conn = get_connection()
        rows = conn.execute("SELECT * FROM clients ORDER BY nom").fetchall()
        n = 0
        for c in rows:
            if not self.row_matches(q, c["nom"], c["prenom"], c["tel"], c["ville"]):
                continue
            rx = conn.execute("SELECT * FROM prescriptions WHERE client_id=? ORDER BY id DESC LIMIT 1",
                               (c["id"],)).fetchone()
            self.tree.insert("", "end", iid=str(c["id"]), tags=("even" if n % 2 == 0 else "odd",),
                              values=(c["nom"], c["prenom"] or "—", c["tel"] or "—", c["ville"] or "—",
                                      f"{c['remise']}%" if c["remise"] else "—",
                                      rx["date"] if rx else "—"))
            n += 1
        conn.close()

        

    def delete_selected(self):
        cid = self._selected_id()
        if not cid:
            messagebox.showinfo("Info", "Sélectionnez un client dans la liste.")
            return
        if not messagebox.askyesno("Confirmer", "Supprimer ce client ?"):
            return
        conn = get_connection()
        row = conn.execute("SELECT nom FROM clients WHERE id=?", (cid,)).fetchone()
        conn.execute("DELETE FROM clients WHERE id=?", (cid,))
        conn.commit()
        conn.close()
        log_action(self.app.current_user["username"], f"Client supprimé : {row['nom']}")
        self.refresh()

    def open_detail(self, client_id=None):
        if not client_id:
            messagebox.showinfo("Info", "Sélectionnez un client.")
            return
        conn = get_connection()
        c = conn.execute("SELECT * FROM clients WHERE id=?", (client_id,)).fetchone()
        rx = conn.execute("SELECT * FROM prescriptions WHERE client_id=? ORDER BY id DESC LIMIT 1",
                           (client_id,)).fetchone()
        conn.close()
        if not c:
            return

        def rxg(key, default="—"):
            try:
                v = rx[key]
                return v if v not in (None, "") else default
            except (KeyError, IndexError, TypeError):
                return default

        def build(body):
            self.info_row(body, "Né(e) le", c["date_naissance"])
            self.info_row(body, "Ville", c["ville"])
            self.info_row(body, "Téléphone", c["tel"])
            self.info_row(body, "Remise", f"{c['remise']}%" if c["remise"] else "—")

            if rx:
                self.section_label(body, "E.F.H")
                efh = " · ".join(lbl for lbl, on in (("VL", rx["vl"]), ("VP", rx["vp"]), ("PG", rx["pg"])) if on) or "—"
                self.info_row(body, "Type", efh)

                self.section_label(body, "ARX")
                self.info_row(body, "OD", rxg("arx_od"))
                self.info_row(body, "OG", rxg("arx_og"))

                self.section_label(body, "AV. Brute")
                self.info_row(body, "OD / OG / ODG",
                              f"{rxg('av_od')} / {rxg('av_og')} / {rxg('av_odg')}")

                self.section_label(body, "Suivi")
                self.info_row(body, "Prescripteur", rxg("prescripteur"))
                self.info_row(body, "Mutuelle", "Oui — " + rxg("mutuelle_texte") if rx["mutuelle_oui"] else "Non")
                self.info_row(body, "Maladie", "Oui — " + rxg("maladie_texte") if rx["maladie_oui"] else "Non")
                if rx["montage_oui"]:
                    montage_display = "Oui" + (" — " + rx["montage_texte"] if rx["montage_texte"] else "")
                else:
                    montage_display = "Non"
                self.info_row(body, "Montage", montage_display)

                self.section_label(body, "Correction")
                table = tk.Frame(body, bg=SURFACE)
                table.pack(fill="x", pady=(0, 4))
                cols = ["", "SPH", "CYL", "AXE", "ADD", "EP", "HT"]
                for i, h in enumerate(cols):
                    tk.Label(table, text=h, font=FONT_MONO_SM, fg=SLATE, bg=SURFACE).grid(row=0, column=i, padx=4)
                    table.grid_columnconfigure(i, weight=1)
                for r, eye in ((1, "od"), (2, "og")):
                    tk.Label(table, text=eye.upper(), font=FONT_BODY_B, fg=TEXT, bg=SURFACE).grid(row=r, column=0)
                    for i, field_key in enumerate(("sph", "cyl", "axe", "add", "ep", "ht"), start=1):
                        tk.Label(table, text=rxg(f"{eye}_{field_key}"), font=FONT_MONO_SM, fg=TEXT,
                                  bg=SURFACE).grid(row=r, column=i, padx=4, pady=2)

                self.section_label(body, "Verres / Firme / Monture")
                self.info_row(body, "Verres", rxg("verres_texte"))
                self.info_row(body, "Firme", rxg("firme_texte"))
                self.info_row(body, "Monture", rxg("monture_texte"))

                self.section_label(body, "Prix")
                self.info_row(body, "Prix / Avance", f"{rxg('prix')} / {rxg('avance')}")
                self.info_row(body, "Jour de livraison", rxg("jour_livraison"))
                self.info_row(body, "R", rxg("r_texte"))
            else:
                self.section_label(body, "Prescription")
                tk.Label(body, text="Aucune prescription enregistrée pour ce client.",
                          font=FONT_MONO_SM, bg=SURFACE, fg=SLATE).pack(anchor="w")

        title = f"{c['nom']} {c['prenom'] or ''}".strip() or "Fiche client"
        self.show_detail_modal(title, build,
                                lambda: self.open_form(client_id),
                                on_history=lambda: self.open_history(client_id),
                                width=480)

    def open_history(self, client_id):
        """Historique complet : chaque prescription est affichée comme une fiche
        complète (ARX, AV. Brute, Suivi, Correction, Verres/Firme/Monture, Prix),
        empilée de la plus récente à la plus ancienne."""
        conn = get_connection()
        c = conn.execute("SELECT * FROM clients WHERE id=?", (client_id,)).fetchone()
        rxs = conn.execute("SELECT * FROM prescriptions WHERE client_id=? ORDER BY id DESC",
                            (client_id,)).fetchall()
        conn.close()
        if not c:
            return

        name = f"{c['nom']} {c['prenom'] or ''}".strip() or "Client"
        modal = ModalForm(self.app, f"Historique — {name}", width=820)
        modal.geometry("820x680")
        modal.center()

        tk.Label(modal.body, text=f"{len(rxs)} fiche(s) enregistrée(s)",
                 font=FONT_MONO_SM, fg=SLATE, bg=SURFACE).pack(anchor="w", pady=(0, 10))

        if not rxs:
            tk.Label(modal.body, text="Aucune fiche enregistrée pour ce client.",
                     font=FONT_BODY, fg=SLATE, bg=SURFACE).pack(anchor="w", pady=20)
            outline_button(modal.footer, "Fermer", modal.destroy, icon="close").pack(side="left", ipady=4)
            return

        def g(rx, key, default="—"):
            try:
                v = rx[key]
                return v if v not in (None, "") else default
            except (KeyError, IndexError, TypeError):
                return default

        def info(parent, label, value):
            row = tk.Frame(parent, bg=SURFACE)
            row.pack(fill="x", pady=2)
            tk.Label(row, text=label.upper(), font=FONT_MONO_SM, bg=SURFACE, fg=SLATE).pack(side="left")
            tk.Label(row, text=value if value not in (None, "") else "—", font=FONT_BODY_B,
                     bg=SURFACE, fg=TEXT, wraplength=460, justify="right").pack(side="right")

        def section(parent, text):
            tk.Label(parent, text=text, font=FONT_MONO_SM, fg=RED, bg=SURFACE).pack(anchor="w", pady=(8, 2))

        for rx in rxs:
            card = tk.Frame(modal.body, bg=SURFACE, highlightbackground=BORDER,
                             highlightthickness=1, padx=16, pady=14)
            card.pack(fill="x", pady=(0, 14))

            # En-tête : date + badges E.F.H
            head = tk.Frame(card, bg=SURFACE)
            head.pack(fill="x", pady=(0, 4))
            tk.Label(head, text=rx["date"] or "—", font=FONT_DISPLAY_SM,
                     bg=SURFACE, fg=TEXT).pack(side="left")
            efh = " · ".join(lbl for lbl, on in (("VL", rx["vl"]), ("VP", rx["vp"]), ("PG", rx["pg"])) if on)
            if efh:
                tk.Label(head, text=efh, font=FONT_MONO_SM, bg=SURFACE, fg=RED).pack(side="right")

            section(card, "ARX")
            info(card, "OD", g(rx, "arx_od"))
            info(card, "OG", g(rx, "arx_og"))

            section(card, "AV. Brute")
            info(card, "OD / OG / ODG", f"{g(rx, 'av_od')} / {g(rx, 'av_og')} / {g(rx, 'av_odg')}")

            section(card, "Suivi")
            info(card, "Prescripteur", g(rx, "prescripteur"))
            info(card, "Date prescription", g(rx, "date_prescription"))
            info(card, "Mutuelle",
                 ("Oui — " + g(rx, "mutuelle_texte")) if rx["mutuelle_oui"] else "Non")
            info(card, "Maladie",
                 ("Oui — " + g(rx, "maladie_texte")) if rx["maladie_oui"] else "Non")
            if rx["montage_oui"]:
                montage_display = "Oui" + (" — " + g(rx, "montage_texte", "") if g(rx, "montage_texte", "") else "")
            else:
                montage_display = "Non"
            info(card, "Montage", montage_display)

            section(card, "Correction")
            tbl = tk.Frame(card, bg=SURFACE)
            tbl.pack(fill="x", pady=(0, 4))
            cols = ["", "SPH", "CYL", "AXE", "ADD", "EP", "HT"]
            for i, h in enumerate(cols):
                tk.Label(tbl, text=h, font=FONT_MONO_SM, fg=SLATE, bg=SURFACE
                          ).grid(row=0, column=i, padx=6)
                tbl.grid_columnconfigure(i, weight=1)
            for r, eye in ((1, "od"), (2, "og")):
                tk.Label(tbl, text=eye.upper(), font=FONT_BODY_B, fg=TEXT,
                          bg=SURFACE).grid(row=r, column=0)
                for i, key in enumerate(("sph", "cyl", "axe", "add", "ep", "ht"), start=1):
                    tk.Label(tbl, text=g(rx, f"{eye}_{key}"), font=FONT_MONO_SM,
                              fg=TEXT, bg=SURFACE).grid(row=r, column=i, padx=6, pady=2)

            section(card, "Verres / Firme / Monture")
            v_txt = (g(rx, "verres_texte") + ("  ✓" if rx["verres_check"] else "")).strip()
            f_txt = (g(rx, "firme_texte") + ("  ✓" if rx["firme_check"] else "")).strip()
            m_txt = (g(rx, "monture_texte") + ("  ✓" if rx["monture_check"] else "")).strip()
            info(card, "Verres", v_txt)
            info(card, "Firme", f_txt)
            info(card, "Monture", m_txt)

            section(card, "Prix")
            info(card, "Prix / Avance", f"{g(rx, 'prix')} / {g(rx, 'avance')}")
            info(card, "Jour de livraison", g(rx, "jour_livraison"))
            info(card, "R", g(rx, "r_texte"))

        outline_button(modal.footer, "Fermer", modal.destroy, icon="close").pack(side="left", ipady=4)

    def open_form(self, client_id=None):
        conn = get_connection()
        c = conn.execute("SELECT * FROM clients WHERE id=?", (client_id,)).fetchone() if client_id else None
        rx = conn.execute("SELECT * FROM prescriptions WHERE client_id=? ORDER BY id DESC LIMIT 1",
                           (client_id,)).fetchone() if client_id else None
        conn.close()
        rx = rx or {}

        def rxg(key, default=""):
            try:
                v = rx[key]
                return v if v is not None else default
            except (KeyError, IndexError, TypeError):
                return default

        modal = ModalForm(self.app, "Modifier la fiche client" if c else "Nouvelle fiche client", width=560)
        body = modal.body

        # ---- identité ----
        row1 = tk.Frame(body, bg=SURFACE); row1.pack(fill="x")
        row1.grid_columnconfigure(0, weight=1); row1.grid_columnconfigure(1, weight=1)
        nom_wrap = tk.Frame(row1, bg=SURFACE); nom_wrap.grid(row=0, column=0, sticky="ew", padx=(0, 6))
        prenom_wrap = tk.Frame(row1, bg=SURFACE); prenom_wrap.grid(row=0, column=1, sticky="ew", padx=(6, 0))
        nom_e = field(nom_wrap, "Nom", c["nom"] if c else "")
        prenom_e = field(prenom_wrap, "Prénom", c["prenom"] if c else "")

        row2 = tk.Frame(body, bg=SURFACE); row2.pack(fill="x")
        row2.grid_columnconfigure(0, weight=1); row2.grid_columnconfigure(1, weight=1)
        naiss_wrap = tk.Frame(row2, bg=SURFACE); naiss_wrap.grid(row=0, column=0, sticky="ew", padx=(0, 6))
        ville_wrap = tk.Frame(row2, bg=SURFACE); ville_wrap.grid(row=0, column=1, sticky="ew", padx=(6, 0))
        naiss_e = field_date(naiss_wrap, "Né(e) le", c["date_naissance"] if c else "")
        ville_e = field(ville_wrap, "Ville", c["ville"] if c else "")

        row3 = tk.Frame(body, bg=SURFACE); row3.pack(fill="x")
        row3.grid_columnconfigure(0, weight=1); row3.grid_columnconfigure(1, weight=1)
        num_wrap = tk.Frame(row3, bg=SURFACE); num_wrap.grid(row=0, column=0, sticky="ew", padx=(0, 6))
        remise_wrap = tk.Frame(row3, bg=SURFACE); remise_wrap.grid(row=0, column=1, sticky="ew", padx=(6, 0))
        num_e = field(num_wrap, "Num", c["tel"] if c else "")
        remise_options = [f"{d}%" for d in DISCOUNT_OPTIONS]
        remise_initial = f"{c['remise']}%" if c and c["remise"] else "0%"
        remise_v = dropdown(remise_wrap, "Remise", remise_options, remise_initial)

        tk.Frame(body, bg=BORDER, height=1).pack(fill="x", pady=(12, 10))

        # ---- E.F.H : VL / VP / PG ----
        tk.Label(body, text="E.F.H", font=FONT_MONO_SM, fg=RED, bg=SURFACE).pack(anchor="w")
        efh_row = tk.Frame(body, bg=SURFACE); efh_row.pack(fill="x", pady=(4, 8))
        vl_var = tk.BooleanVar(value=bool(rxg("vl", 0)))
        vp_var = tk.BooleanVar(value=bool(rxg("vp", 0)))
        pg_var = tk.BooleanVar(value=bool(rxg("pg", 0)))
        for label, var in (("VL", vl_var), ("VP", vp_var), ("PG", pg_var)):
            tk.Checkbutton(efh_row, text=label, variable=var, font=FONT_BODY_B, bg=SURFACE, fg=TEXT,
                            activebackground=SURFACE, activeforeground=TEXT, selectcolor=SURFACE
                            ).pack(side="left", padx=(0, 16))

        tk.Label(body, text="ARX", font=FONT_MONO_SM, fg=SLATE, bg=SURFACE).pack(anchor="w", pady=(6, 0))
        arx_row = tk.Frame(body, bg=SURFACE); arx_row.pack(fill="x")
        arx_row.grid_columnconfigure(0, weight=1); arx_row.grid_columnconfigure(1, weight=1)
        arx_od_wrap = tk.Frame(arx_row, bg=SURFACE); arx_od_wrap.grid(row=0, column=0, sticky="ew", padx=(0, 6))
        arx_og_wrap = tk.Frame(arx_row, bg=SURFACE); arx_og_wrap.grid(row=0, column=1, sticky="ew", padx=(6, 0))
        arx_od_e = field(arx_od_wrap, "OD", rxg("arx_od"))
        arx_og_e = field(arx_og_wrap, "OG", rxg("arx_og"))

        tk.Label(body, text="AV. BRUTE", font=FONT_MONO_SM, fg=SLATE, bg=SURFACE).pack(anchor="w", pady=(6, 0))
        av_row = tk.Frame(body, bg=SURFACE); av_row.pack(fill="x")
        for i in range(3):
            av_row.grid_columnconfigure(i, weight=1)
        av_od_wrap = tk.Frame(av_row, bg=SURFACE); av_od_wrap.grid(row=0, column=0, sticky="ew", padx=(0, 4))
        av_og_wrap = tk.Frame(av_row, bg=SURFACE); av_og_wrap.grid(row=0, column=1, sticky="ew", padx=4)
        av_odg_wrap = tk.Frame(av_row, bg=SURFACE); av_odg_wrap.grid(row=0, column=2, sticky="ew", padx=(4, 0))
        av_od_e = field(av_od_wrap, "OD /10", rxg("av_od"))
        av_og_e = field(av_og_wrap, "OG /10", rxg("av_og"))
        av_odg_e = field(av_odg_wrap, "ODG /10", rxg("av_odg"))

        tk.Frame(body, bg=BORDER, height=1).pack(fill="x", pady=(12, 10))

        # ---- Prescripteur / Mutuelle / Maladie / Montage ----
        presc_row = tk.Frame(body, bg=SURFACE); presc_row.pack(fill="x")
        presc_row.grid_columnconfigure(0, weight=2); presc_row.grid_columnconfigure(1, weight=1)
        presc_wrap = tk.Frame(presc_row, bg=SURFACE); presc_wrap.grid(row=0, column=0, sticky="ew", padx=(0, 6))
        presc_date_wrap = tk.Frame(presc_row, bg=SURFACE); presc_date_wrap.grid(row=0, column=1, sticky="ew", padx=(6, 0))
        prescripteur_e = field(presc_wrap, "Prescripteur", rxg("prescripteur"))
        presc_date_e = field_date(presc_date_wrap, "Le", rxg("date_prescription", today_iso()))

        mutuelle_var = oui_non(body, "Mutuelle", bool(rxg("mutuelle_oui", 0)))
        mutuelle_txt_e = field(body, "Mutuelle — précisions", rxg("mutuelle_texte"))
        maladie_var = oui_non(body, "Maladie", bool(rxg("maladie_oui", 0)))
        maladie_txt_e = field(body, "Maladie — précisions", rxg("maladie_texte"))
        montage_var = oui_non(body, "Montage", bool(rxg("montage_oui", 0)))
        montage_txt_e = field(body, "Montage — précisions", rxg("montage_texte"))

        tk.Frame(body, bg=BORDER, height=1).pack(fill="x", pady=(12, 10))

        # ---- Tableau SPH / CYL / AXE / ADD / EP / HT ----
        tk.Label(body, text="CORRECTION", font=FONT_MONO_SM, fg=RED, bg=SURFACE).pack(anchor="w", pady=(0, 6))
        table = tk.Frame(body, bg=SURFACE)
        table.pack(fill="x", pady=(0, 10))
        cols = ["", "SPH", "CYL", "AXE", "ADD", "EP", "HT"]
        for col_i, htext in enumerate(cols):
            tk.Label(table, text=htext, font=FONT_MONO_SM, fg=SLATE, bg=SURFACE
                      ).grid(row=0, column=col_i, padx=2, pady=(0, 4))
            table.grid_columnconfigure(col_i, weight=1)
        tk.Label(table, text="OD", font=FONT_BODY_B, fg=TEXT, bg=SURFACE).grid(row=1, column=0)
        tk.Label(table, text="OG", font=FONT_BODY_B, fg=TEXT, bg=SURFACE).grid(row=2, column=0)
        od_sph = mini_entry(table, 1, 1, rxg("od_sph")); od_cyl = mini_entry(table, 1, 2, rxg("od_cyl"))
        od_axe = mini_entry(table, 1, 3, rxg("od_axe")); od_add = mini_entry(table, 1, 4, rxg("od_add"))
        od_ep = mini_entry(table, 1, 5, rxg("od_ep")); od_ht = mini_entry(table, 1, 6, rxg("od_ht"))
        og_sph = mini_entry(table, 2, 1, rxg("og_sph")); og_cyl = mini_entry(table, 2, 2, rxg("og_cyl"))
        og_axe = mini_entry(table, 2, 3, rxg("og_axe")); og_add = mini_entry(table, 2, 4, rxg("og_add"))
        og_ep = mini_entry(table, 2, 5, rxg("og_ep")); og_ht = mini_entry(table, 2, 6, rxg("og_ht"))

        # ---- Verres / Firme / Monture ----
        def text_with_check(label_text, initial_text, initial_check):
            row = tk.Frame(body, bg=SURFACE); row.pack(fill="x")
            row.grid_columnconfigure(0, weight=1)
            wrap = tk.Frame(row, bg=SURFACE); wrap.grid(row=0, column=0, sticky="ew")
            e = field(wrap, label_text, initial_text)
            cvar = tk.BooleanVar(value=bool(initial_check))
            tk.Checkbutton(row, variable=cvar, bg=SURFACE, activebackground=SURFACE,
                            selectcolor=SURFACE).grid(row=0, column=1, padx=(8, 0), pady=(18, 0))
            return e, cvar

        verres_e, verres_c = text_with_check("Verres", rxg("verres_texte"), rxg("verres_check", 0))
        firme_e, firme_c = text_with_check("Firme", rxg("firme_texte"), rxg("firme_check", 0))
        monture_e, monture_c = text_with_check("Monture", rxg("monture_texte"), rxg("monture_check", 0))

        row_prix = tk.Frame(body, bg=SURFACE); row_prix.pack(fill="x")
        row_prix.grid_columnconfigure(0, weight=1); row_prix.grid_columnconfigure(1, weight=1)
        prix_wrap = tk.Frame(row_prix, bg=SURFACE); prix_wrap.grid(row=0, column=0, sticky="ew", padx=(0, 6))
        avance_wrap = tk.Frame(row_prix, bg=SURFACE); avance_wrap.grid(row=0, column=1, sticky="ew", padx=(6, 0))
        prix_e = field(prix_wrap, "Prix", rxg("prix"))
        avance_e = field(avance_wrap, "Avance", rxg("avance"))

        tk.Label(body, text="JOUR DE LIVRAISON", font=FONT_MONO_SM, fg=SLATE, bg=SURFACE).pack(anchor="w", pady=(10, 4))
        jour_var = tk.StringVar(value=rxg("jour_livraison"))
        days_row = tk.Frame(body, bg=SURFACE); days_row.pack(fill="x")
        for j in JOURS_SEMAINE:
            tk.Radiobutton(days_row, text=j, variable=jour_var, value=j, indicatoron=0,
                            font=FONT_MONO_SM, bg=SURFACE, fg=TEXT, selectcolor=RED,
                            activebackground=SURFACE, width=4, relief="solid", borderwidth=1
                            ).pack(side="left", padx=(0, 4))

        r_e = field(body, "R", rxg("r_texte"))

        def save():
            nom = nom_e.get().strip()
            if not nom:
                messagebox.showwarning("Champ requis", "Le nom est obligatoire.")
                return
            remise_val = int(remise_v.get().replace("%", "") or 0)
            conn = get_connection()
            if c:
                conn.execute("""UPDATE clients SET nom=?, prenom=?, tel=?, date_naissance=?, ville=?, remise=?
                                 WHERE id=?""",
                             (nom, prenom_e.get(), num_e.get(), naiss_e.get(), ville_e.get(), remise_val, c["id"]))
                cid = c["id"]
                action = f"Client modifié : {nom}"
            else:
                cid = conn.execute("""INSERT INTO clients (nom, prenom, tel, date_naissance, ville, remise)
                                       VALUES (?,?,?,?,?,?)""",
                                    (nom, prenom_e.get(), num_e.get(), naiss_e.get(), ville_e.get(), remise_val)).lastrowid
                action = f"Client créé : {nom}"

            conn.execute("""INSERT INTO prescriptions (
                client_id, date, vl, vp, pg, arx_od, arx_og, av_od, av_og, av_odg,
                prescripteur, date_prescription, mutuelle_oui, mutuelle_texte,
                maladie_oui, maladie_texte, montage_oui, montage_texte,
                od_sph, od_cyl, od_axe, od_add, od_ep, od_ht,
                og_sph, og_cyl, og_axe, og_add, og_ep, og_ht,
                verres_texte, verres_check, firme_texte, firme_check, monture_texte, monture_check,
                prix, avance, jour_livraison, r_texte
            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", (
                cid, today_iso(), int(vl_var.get()), int(vp_var.get()), int(pg_var.get()),
                arx_od_e.get(), arx_og_e.get(), av_od_e.get(), av_og_e.get(), av_odg_e.get(),
                prescripteur_e.get(), presc_date_e.get(),
                1 if mutuelle_var.get() == "oui" else 0, mutuelle_txt_e.get(),
                1 if maladie_var.get() == "oui" else 0, maladie_txt_e.get(),
                1 if montage_var.get() == "oui" else 0, montage_txt_e.get(),
                od_sph.get(), od_cyl.get(), od_axe.get(), od_add.get(), od_ep.get(), od_ht.get(),
                og_sph.get(), og_cyl.get(), og_axe.get(), og_add.get(), og_ep.get(), og_ht.get(),
                verres_e.get(), int(verres_c.get()), firme_e.get(), int(firme_c.get()),
                monture_e.get(), int(monture_c.get()),
                prix_e.get(), avance_e.get(), jour_var.get(), r_e.get(),
            ))
            conn.commit()
            conn.close()
            log_action(self.app.current_user["username"], action)
            modal.destroy()
            self.refresh()

        modal.buttons(save)


# ================================================================== APPOINTMENTS
class AppointmentsPage(BasePage):
    def __init__(self, parent, app):
        super().__init__(parent, app)
        self.header("Agenda", "Rendez-vous", "+ Nouveau RDV", self.open_form)
        self.search_var = self.make_search_box("Rechercher un rendez-vous…")
        self.tree = self.make_table(["Date", "Heure", "Client", "Service", "Statut"])
        self.tree.bind("<Double-1>", lambda e: self.open_form(self._selected_id()))
        btns = tk.Frame(self, bg=CONTENT_BG)
        btns.pack(fill="x", pady=(8, 0))
        outline_button(btns, "Modifier la sélection", lambda: self.open_form(self._selected_id()), icon="edit").pack(side="left")
        danger_button(btns, "Supprimer", self.delete_selected).pack(side="left", padx=(0, 12))
        outline_button(btns, "Envoyer rappel WhatsApp", self.send_whatsapp_reminder, icon="bell").pack(side="left")

    def send_whatsapp_reminder(self):
        aid = self._selected_id()
        if not aid:
            messagebox.showinfo("Info", "Sélectionnez un rendez-vous.")
            return
        conn = get_connection()
        a = conn.execute("SELECT * FROM appointments WHERE id=?", (aid,)).fetchone()
        conn.close()
        phone = self.app.find_client_phone(a["client_nom"])
        if not phone:
            messagebox.showwarning("Pas de numéro",
                                    f"Aucun numéro de téléphone trouvé pour « {a['client_nom']} » "
                                    "dans sa fiche client.")
            return
        msg = (f"Bonjour {a['client_nom']}, un rappel de votre rendez-vous chez OPTILUX "
               f"le {a['date']} à {a['heure']} ({a['service']}). À bientôt !")
        sent = send_whatsapp(phone, msg)
        if sent:
            log_action(self.app.current_user["username"], f"Rappel WhatsApp envoyé : {a['client_nom']}")
            messagebox.showinfo("En cours d'envoi",
                                 "WhatsApp Web va s'ouvrir dans votre navigateur pour envoyer le message.\n\n"
                                 "Assurez-vous d'être déjà connecté à WhatsApp Web (une seule fois, par QR code).")

    def _selected_id(self):
        sel = self.tree.selection()
        return int(sel[0]) if sel else None

    def refresh(self):
        for i in self.tree.get_children():
            self.tree.delete(i)
        q = self.search_var.get()
        conn = get_connection()
        rows = conn.execute("SELECT * FROM appointments ORDER BY date, heure").fetchall()
        conn.close()
        n = 0
        for a in rows:
            if not self.row_matches(q, a["client_nom"], a["service"], a["statut"]):
                continue
            self.tree.insert("", "end", iid=str(a["id"]), tags=("even" if n % 2 == 0 else "odd",),
                              values=(a["date"], a["heure"], a["client_nom"], a["service"], a["statut"]))
            n += 1

    def delete_selected(self):
        aid = self._selected_id()
        if not aid:
            messagebox.showinfo("Info", "Sélectionnez un rendez-vous.")
            return
        if not messagebox.askyesno("Confirmer", "Supprimer ce rendez-vous ?"):
            return
        conn = get_connection()
        row = conn.execute("SELECT client_nom FROM appointments WHERE id=?", (aid,)).fetchone()
        conn.execute("DELETE FROM appointments WHERE id=?", (aid,))
        conn.commit(); conn.close()
        log_action(self.app.current_user["username"], f"RDV supprimé : {row['client_nom']}")
        self.refresh()

    def open_form(self, apt_id=None):
        conn = get_connection()
        a = conn.execute("SELECT * FROM appointments WHERE id=?", (apt_id,)).fetchone() if apt_id else None
        client_names = [r["nom"] for r in conn.execute("SELECT nom FROM clients ORDER BY nom").fetchall()]
        conn.close()

        modal = ModalForm(self.app, "Modifier rendez-vous" if a else "Nouveau rendez-vous", width=400)
        body = modal.body
        tk.Label(body, text="CLIENT", font=FONT_MONO_SM, fg=SLATE, bg=SURFACE, anchor="w").pack(fill="x", pady=(8, 2))
        client_var = tk.StringVar(value=a["client_nom"] if a else "")
        cb = ttk.Combobox(body, textvariable=client_var, values=client_names, font=FONT_BODY, style="Themed.TCombobox")
        cb.pack(fill="x", ipady=2)

        date_e = field_date(body, "Date (AAAA-MM-JJ)", a["date"] if a else today_iso())
        heure_e = field(body, "Heure (HH:MM)", a["heure"] if a else "10:00")
        service_v = dropdown(body, "Service", SERVICES_RDV, a["service"] if a else SERVICES_RDV[0])
        statut_v = dropdown(body, "Statut", STATUTS_RDV, a["statut"] if a else STATUTS_RDV[0])

        def save():
            nom = client_var.get().strip()
            if not nom:
                messagebox.showwarning("Champ requis", "Le client est obligatoire.")
                return
            conn = get_connection()
            if a:
                conn.execute("UPDATE appointments SET client_nom=?, date=?, heure=?, service=?, statut=? WHERE id=?",
                             (nom, date_e.get(), heure_e.get(), service_v.get(), statut_v.get(), a["id"]))
                action = f"RDV modifié : {nom}"
            else:
                conn.execute("INSERT INTO appointments (client_nom,date,heure,service,statut) VALUES (?,?,?,?,?)",
                             (nom, date_e.get(), heure_e.get(), service_v.get(), statut_v.get()))
                action = f"RDV créé : {nom}"
            conn.commit(); conn.close()
            log_action(self.app.current_user["username"], action)
            modal.destroy()
            self.refresh()

        modal.buttons(save)


# ================================================================== STOCK
# ================================================================== STOCK
class StockPage(BasePage):
    def __init__(self, parent, app):
        super().__init__(parent, app)
        self.header("Inventaire", "Stock", "+ Nouvel article", self.open_form)
        self.search_var = self.make_search_box("Rechercher un article…")
        self.tree = self.make_table(["Photo", "Article", "Catégorie", "Marque", "Qté", "Prix vente", "Statut"])
        self.tree.bind("<Double-1>", lambda e: self.open_detail(self._selected_id()))
        btns = tk.Frame(self, bg=CONTENT_BG)
        btns.pack(fill="x", pady=(8, 0))
        outline_button(btns, "Voir la fiche", lambda: self.open_detail(self._selected_id()), icon="search").pack(side="left")
        danger_button(btns, "Supprimer", self.delete_selected).pack(side="left", padx=(0, 12))

    def _selected_id(self):
        sel = self.tree.selection()
        return int(sel[0]) if sel else None

    def refresh(self):
        for i in self.tree.get_children():
            self.tree.delete(i)
        q = self.search_var.get()
        conn = get_connection()
        rows = conn.execute("SELECT * FROM stock ORDER BY nom").fetchall()
        conn.close()
        n = 0
        for s in rows:
            if not self.row_matches(q, s["nom"], s["categorie"], s["marque"]):
                continue
            low = s["qte"] <= s["seuil"]
            iid = str(s["id"])
            tag = "low" if low else ("even" if n % 2 == 0 else "odd")
            self.tree.insert("", "end", iid=iid, tags=(tag,),
                              values=("Oui" if s["image_path"] else "—", s["nom"], s["categorie"], s["marque"], s["qte"],
                                      money(s["prix_vente"]), "STOCK BAS" if low else "OK"))
            n += 1

    def delete_selected(self):
        sid = self._selected_id()
        if not sid:
            messagebox.showinfo("Info", "Sélectionnez un article.")
            return
        if not messagebox.askyesno("Confirmer", "Supprimer cet article ?"):
            return
        conn = get_connection()
        row = conn.execute("SELECT nom FROM stock WHERE id=?", (sid,)).fetchone()
        conn.execute("DELETE FROM stock WHERE id=?", (sid,))
        conn.commit(); conn.close()
        log_action(self.app.current_user["username"], f"Article supprimé : {row['nom']}")
        self.refresh()

    def open_detail(self, stock_id=None):
        if not stock_id:
            messagebox.showinfo("Info", "Sélectionnez un article.")
            return
        conn = get_connection()
        s = conn.execute("SELECT * FROM stock WHERE id=?", (stock_id,)).fetchone()
        conn.close()
        if not s:
            return

        def build(body):
            img_wrap = tk.Frame(body, bg=SURFACE)
            img_wrap.pack(fill="x", pady=(0, 14))
            full = os.path.join(PRODUCT_IMAGES_DIR, s["image_path"]) if s["image_path"] else None
            photo_ref = None
            if full and os.path.exists(full):
                try:
                    from PIL import Image, ImageTk
                    im = Image.open(full).convert("RGB")
                    im.thumbnail((320, 320), Image.LANCZOS)
                    photo_ref = ImageTk.PhotoImage(im)
                except Exception:
                    photo_ref = None
            img_lbl = tk.Label(img_wrap, bg=SURFACE)
            if photo_ref:
                img_lbl.configure(image=photo_ref)
                img_lbl._icon_ref = photo_ref
            else:
                img_lbl.configure(text="Aucune photo", fg=SLATE, font=FONT_MONO_SM, width=26, height=10,
                                    relief="solid", borderwidth=1, bg=SURFACE)
            img_lbl.pack(anchor="center", pady=(0, 4))

            low = s["qte"] <= s["seuil"]
            self.info_row(body, "Statut", "STOCK BAS" if low else "OK", value_color=RED if low else "#2E7D46")
            self.info_row(body, "Catégorie", s["categorie"])
            self.info_row(body, "Marque", s["marque"])
            self.info_row(body, "Quantité", str(s["qte"]))
            self.info_row(body, "Prix achat", money(s["prix_achat"]))
            self.info_row(body, "Prix vente", money(s["prix_vente"]))
            self.info_row(body, "Seuil d'alerte", str(s["seuil"]))
            if s["categorie"] == "Lentilles":
                self.section_label(body, "Lentilles")
                self.info_row(body, "Type", s["lentille_type"])
                self.info_row(body, "Durée de port", s["lentille_duree"])

        self.show_detail_modal(s["nom"], build, lambda: self.open_form(stock_id), width=440)

    def open_form(self, stock_id=None):
        conn = get_connection()
        s = conn.execute("SELECT * FROM stock WHERE id=?", (stock_id,)).fetchone() if stock_id else None
        conn.close()

        cats = get_types("stock_categorie") or ["Autre"]
        marques = get_types("stock_marque") or ["Autre"]

        modal = ModalForm(self.app, "Modifier article" if s else "Nouvel article", width=420)
        body = modal.body
        nom_e = field(body, "Nom", s["nom"] if s else "")
        cat_v = dropdown(body, "Catégorie", cats, s["categorie"] if s else cats[0])
        marque_v = dropdown(body, "Marque", marques, s["marque"] if s else marques[0])

        row = tk.Frame(body, bg=SURFACE); row.pack(fill="x")
        qte_e = field(row, "Quantité", str(s["qte"]) if s else "0")
        row2 = tk.Frame(body, bg=SURFACE); row2.pack(fill="x")
        achat_e = field(row2, "Prix achat", str(s["prix_achat"]) if s else "0")
        vente_e = field(row2, "Prix vente", str(s["prix_vente"]) if s else "0")
        seuil_e = field(body, "Seuil d'alerte", str(s["seuil"]) if s else "3")

        # ---- champs lentilles (affichés seulement si catégorie = Lentilles) ----
        lentille_frame = tk.Frame(body, bg=SURFACE)
        lentille_types = get_types("lentille_type") or ["Autre"]
        lentille_durees = get_types("lentille_duree") or ["Autre"]
        lentille_type_v = dropdown(lentille_frame, "Type de lentille",
                                    lentille_types, s["lentille_type"] if s and s["lentille_type"] else lentille_types[0])
        lentille_duree_v = dropdown(lentille_frame, "Durée de port",
                                     lentille_durees, s["lentille_duree"] if s and s["lentille_duree"] else lentille_durees[0])

        def sync_lentille_visibility(*_a):
            if cat_v.get() == "Lentilles":
                lentille_frame.pack(fill="x")
            else:
                lentille_frame.pack_forget()
        cat_v.trace_add("write", sync_lentille_visibility)
        sync_lentille_visibility()

        # ---- photo produit ----
        tk.Label(body, text="PHOTO", font=FONT_MONO_SM, fg=SLATE, bg=SURFACE, anchor="w").pack(fill="x", pady=(10, 2))
        photo_row = tk.Frame(body, bg=SURFACE); photo_row.pack(fill="x")
        preview_lbl = tk.Label(photo_row, bg=SURFACE, text="Aucune image", fg=SLATE, font=FONT_MONO_SM,
                                width=12, height=6, relief="solid", borderwidth=1)
        preview_lbl.pack(side="left")
        state = {"image_path": s["image_path"] if s else None, "preview_ref": None}

        def refresh_preview():
            path = state["image_path"]
            full = os.path.join(PRODUCT_IMAGES_DIR, path) if path else None
            if full and os.path.exists(full):
                try:
                    from PIL import Image, ImageTk
                    im = Image.open(full).convert("RGB")
                    im.thumbnail((96, 96), Image.LANCZOS)
                    photo = ImageTk.PhotoImage(im)
                    state["preview_ref"] = photo
                    preview_lbl.configure(image=photo, text="")
                except Exception:
                    preview_lbl.configure(image="", text="Image")
            else:
                preview_lbl.configure(image="", text="Aucune image")

        def choose_image():
            path = filedialog.askopenfilename(
                title="Choisir une photo",
                filetypes=[("Images", "*.png *.jpg *.jpeg *.webp *.bmp")])
            if not path:
                return
            try:
                from PIL import Image
                os.makedirs(PRODUCT_IMAGES_DIR, exist_ok=True)
                fname = f"{uuid.uuid4().hex}.png"
                im = Image.open(path).convert("RGB")
                im.thumbnail((600, 600), Image.LANCZOS)
                im.save(os.path.join(PRODUCT_IMAGES_DIR, fname))
                state["image_path"] = fname
                refresh_preview()
            except Exception as e:
                messagebox.showerror("Erreur", f"Impossible de charger l'image :\n{e}")

        def remove_image():
            state["image_path"] = None
            refresh_preview()

        photo_btns = tk.Frame(photo_row, bg=SURFACE); photo_btns.pack(side="left", padx=10)
        outline_button(photo_btns, "Choisir…", choose_image).pack(fill="x", pady=(0, 6))
        danger_button(photo_btns, "Retirer", remove_image, icon=None).pack(fill="x")
        refresh_preview()

        def save():
            nom = nom_e.get().strip()
            if not nom:
                messagebox.showwarning("Champ requis", "Le nom est obligatoire.")
                return
            try:
                qte = int(float(qte_e.get() or 0))
                achat = float(achat_e.get() or 0)
                vente = float(vente_e.get() or 0)
                seuil = int(float(seuil_e.get() or 0))
            except ValueError:
                messagebox.showwarning("Valeur invalide", "Vérifiez les champs numériques.")
                return
            is_lentille = cat_v.get() == "Lentilles"
            l_type = lentille_type_v.get() if is_lentille else None
            l_duree = lentille_duree_v.get() if is_lentille else None
            conn = get_connection()
            if s:
                conn.execute("""UPDATE stock SET nom=?, categorie=?, marque=?, qte=?, prix_achat=?, prix_vente=?,
                                 seuil=?, image_path=?, lentille_type=?, lentille_duree=? WHERE id=?""",
                             (nom, cat_v.get(), marque_v.get(), qte, achat, vente, seuil,
                              state["image_path"], l_type, l_duree, s["id"]))
                action = f"Article modifié : {nom}"
            else:
                conn.execute("""INSERT INTO stock (nom,categorie,marque,qte,prix_achat,prix_vente,seuil,
                                 image_path,lentille_type,lentille_duree) VALUES (?,?,?,?,?,?,?,?,?,?)""",
                             (nom, cat_v.get(), marque_v.get(), qte, achat, vente, seuil,
                              state["image_path"], l_type, l_duree))
                action = f"Article créé : {nom}"
            conn.commit(); conn.close()
            log_action(self.app.current_user["username"], action)
            if qte <= seuil:
                notify("Stock faible", f"{nom} — {qte} restant(s)", icon_path=ICON_ICO)
            modal.destroy()
            self.refresh()

        modal.buttons(save)


# ================================================================== SALES
class SalesPage(BasePage):
    def __init__(self, parent, app):
        super().__init__(parent, app)
        self.header("Registre", "Ventes", "+ Nouvelle vente", self.open_form)
        self.search_var = self.make_search_box("Rechercher une vente (client, paiement)…")

        self.totals_row = tk.Frame(self, bg=CONTENT_BG)
        self.totals_row.pack(fill="x", pady=(0, 10))

        self.tree = self.make_table(["Date", "Client", "Vente", "Avance", "Reste", "Remise", "Paiement"])
        self.tree.bind("<Double-1>", lambda e: self.open_form(self._selected_id()))
        btns = tk.Frame(self, bg=CONTENT_BG)
        btns.pack(fill="x", pady=(8, 0))
        outline_button(btns, "Modifier la sélection", lambda: self.open_form(self._selected_id()), icon="edit").pack(side="left")
        danger_button(btns, "Supprimer", self.delete_selected).pack(side="left", padx=(0, 12))
        outline_button(btns, "Facture PDF", self.print_facture, icon="document").pack(side="left", padx=(0, 12))
        outline_button(btns, "Facture Word", self.print_facture_docx, icon="document").pack(side="left")

    def _selected_id(self):
        sel = self.tree.selection()
        return int(sel[0]) if sel else None

    def refresh(self):
        for i in self.tree.get_children():
            self.tree.delete(i)
        for w in self.totals_row.winfo_children():
            w.destroy()
        q = self.search_var.get()
        conn = get_connection()
        rows = conn.execute("""SELECT sales.*, clients.nom as client_nom FROM sales
                                LEFT JOIN clients ON clients.id = sales.client_id
                                ORDER BY date DESC""").fetchall()
        conn.close()

        total_tpe, total_especes, n = 0.0, 0.0, 0
        for s in rows:
            if not self.row_matches(q, s["client_nom"], s["paiement"]):
                continue
            reste = (s["vente"] or 0) - (s["avance"] or 0)
            self.tree.insert("", "end", iid=str(s["id"]), tags=("even" if n % 2 == 0 else "odd",),
                              values=(s["date"], s["client_nom"] or "—", money(s["vente"]), money(s["avance"]),
                                      money(reste), f"{s['remise']}%" if s["remise"] else "—",
                                      s["paiement"] or "Espèces"))
            n += 1
            if s["paiement"] == "TPE":
                total_tpe += s["avance"] or 0
            else:
                total_especes += s["avance"] or 0

        self._total_card("Total Espèces", money(total_especes), "sales")
        self._total_card("Total TPE", money(total_tpe), "wallet")
        self._total_card("Total encaissé", money(total_tpe + total_especes), "growth", accent=True)

    def _total_card(self, label, value, icon, accent=False):
        card = tk.Frame(self.totals_row, bg=RED if accent else CARD_DARK, padx=14, pady=10)
        card.pack(side="left", fill="both", expand=True, padx=4)
        top = tk.Frame(card, bg=RED if accent else CARD_DARK); top.pack(fill="x")
        tk.Label(top, text=label.upper(), font=FONT_MONO_SM, bg=RED if accent else CARD_DARK,
                  fg=WHITE if accent else SLATE).pack(side="left")
        icon_photo = get_icon(icon, WHITE, size=14)
        if icon_photo:
            il = tk.Label(top, image=icon_photo, bg=RED if accent else CARD_DARK)
            il._icon_ref = icon_photo
            il.pack(side="right")
        tk.Label(card, text=value, font=FONT_DISPLAY_SM, bg=RED if accent else CARD_DARK, fg=WHITE).pack(anchor="w", pady=(4, 0))

    def delete_selected(self):
        sid = self._selected_id()
        if not sid:
            messagebox.showinfo("Info", "Sélectionnez une vente.")
            return
        if not messagebox.askyesno("Confirmer", "Supprimer cette vente ?"):
            return
        conn = get_connection()
        conn.execute("DELETE FROM sales WHERE id=?", (sid,))
        conn.commit(); conn.close()
        log_action(self.app.current_user["username"], "Vente supprimée")
        self.refresh()

    def open_form(self, sale_id=None):
        conn = get_connection()
        s = conn.execute("SELECT * FROM sales WHERE id=?", (sale_id,)).fetchone() if sale_id else None
        clients = conn.execute("SELECT * FROM clients ORDER BY nom").fetchall()
        conn.close()
        client_map = {f"{c['nom']} {c['prenom'] or ''}".strip(): c["id"] for c in clients}
        client_remise = {c["id"]: (c["remise"] or 0) for c in clients}
        current_name = ""
        if s and s["client_id"]:
            for name, cid in client_map.items():
                if cid == s["client_id"]:
                    current_name = name

        modal = ModalForm(self.app, "Modifier vente" if s else "Nouvelle vente", width=400)
        body = modal.body
        tk.Label(body, text="CLIENT", font=FONT_MONO_SM, fg=SLATE, bg=SURFACE, anchor="w").pack(fill="x", pady=(8, 2))
        client_var = tk.StringVar(value=current_name)
        cb = ttk.Combobox(body, textvariable=client_var, values=list(client_map.keys()),
                           state="readonly", font=FONT_BODY, style="Themed.TCombobox")
        cb.pack(fill="x", ipady=2)

        date_e = field(body, "Date (AAAA-MM-JJ)", s["date"] if s else today_iso())
        row = tk.Frame(body, bg=SURFACE); row.pack(fill="x")
        monture_e = field(row, "Monture (DH)", str(s["monture"]) if s else "0")
        verres_e = field(row, "Verres / LSH (DH)", str(s["verres"]) if s else "0")
        row2 = tk.Frame(body, bg=SURFACE); row2.pack(fill="x")
        vente_e = field(row2, "Vente totale (DH)", str(s["vente"]) if s else "0")
        avance_e = field(row2, "Avance (DH)", str(s["avance"]) if s else "0")
        mutuelle_e = field(body, "Mutuelle (DH, optionnel)", str(s["mutuelle"]) if s else "0")

        row3 = tk.Frame(body, bg=SURFACE); row3.pack(fill="x")
        row3.grid_columnconfigure(0, weight=1); row3.grid_columnconfigure(1, weight=1)
        paiement_wrap = tk.Frame(row3, bg=SURFACE); paiement_wrap.grid(row=0, column=0, sticky="ew", padx=(0, 6))
        remise_wrap = tk.Frame(row3, bg=SURFACE); remise_wrap.grid(row=0, column=1, sticky="ew", padx=(6, 0))
        paiement_v = dropdown(paiement_wrap, "Paiement", MODES_PAIEMENT, s["paiement"] if s else "Espèces")
        remise_options = [f"{d}%" for d in DISCOUNT_OPTIONS]
        default_remise = f"{s['remise']}%" if s and s["remise"] else "0%"
        remise_v = dropdown(remise_wrap, "Remise", remise_options, default_remise)

        def on_client_change(*_a):
            if not s and client_var.get() in client_map:  # ne préremplit que pour une nouvelle vente
                cid = client_map[client_var.get()]
                r = client_remise.get(cid, 0)
                if r:
                    remise_v.set(f"{r}%")
        client_var.trace_add("write", on_client_change)

        def save():
            if not client_var.get():
                messagebox.showwarning("Champ requis", "Sélectionnez un client.")
                return
            try:
                monture = float(monture_e.get() or 0)
                verres = float(verres_e.get() or 0)
                vente = float(vente_e.get() or 0)
                avance = float(avance_e.get() or 0)
                mutuelle = float(mutuelle_e.get() or 0)
            except ValueError:
                messagebox.showwarning("Valeur invalide", "Vérifiez les champs numériques.")
                return
            remise_val = int(remise_v.get().replace("%", "") or 0)
            cid = client_map[client_var.get()]
            conn = get_connection()
            if s:
                conn.execute("""UPDATE sales SET client_id=?, date=?, monture=?, verres=?, vente=?, avance=?,
                                 mutuelle=?, paiement=?, remise=? WHERE id=?""",
                             (cid, date_e.get(), monture, verres, vente, avance, mutuelle,
                              paiement_v.get(), remise_val, s["id"]))
                action = f"Vente modifiée : {client_var.get()}"
            else:
                conn.execute("""INSERT INTO sales (client_id,date,monture,verres,vente,avance,mutuelle,paiement,remise)
                                 VALUES (?,?,?,?,?,?,?,?,?)""",
                             (cid, date_e.get(), monture, verres, vente, avance, mutuelle,
                              paiement_v.get(), remise_val))
                action = f"Vente créée : {client_var.get()} ({money(vente)}, {paiement_v.get()})"
            conn.commit(); conn.close()
            log_action(self.app.current_user["username"], action)
            modal.destroy()
            self.refresh()

        modal.buttons(save)

    def print_facture(self):
        self._generate_and_store_facture(generate_facture_pdf, "PDF", "Impossible de générer le PDF")

    def print_facture_docx(self):
        self._generate_and_store_facture(generate_facture_docx, "Word", "Impossible de générer le document Word")

    def _generate_and_store_facture(self, generator_fn, format_label, error_prefix):
        sid = self._selected_id()
        if not sid:
            messagebox.showinfo("Info", "Sélectionnez une vente pour générer la facture.")
            return                                                                                                                                                                                           
        conn = get_connection()
        s = conn.execute("SELECT * FROM sales WHERE id=?", (sid,)).fetchone()
        c = conn.execute("SELECT * FROM clients WHERE id=?", (s["client_id"],)).fetchone() if s["client_id"] else None
        rx = conn.execute("SELECT * FROM prescriptions WHERE client_id=? ORDER BY id DESC LIMIT 1",
                           (s["client_id"],)).fetchone() if c else None
        conn.close()
        try:
            path = generator_fn(s, c, rx)
        except Exception as e:
            messagebox.showerror("Erreur", f"{error_prefix} :\n{e}")
            return
        client_label = f"{c['nom']} {c['prenom'] or ''}".strip() if c else "Client"
        conn = get_connection()
        existing = conn.execute("SELECT id FROM factures WHERE sale_id=? AND fichier_path=?",
                                 (sid, path)).fetchone()
        if existing:
            conn.execute("UPDATE factures SET titre=?, date=?, montant=? WHERE id=?",
                         (f"{client_label} ({format_label})", today_iso(), s["vente"], existing["id"]))
        else:
            conn.execute("""INSERT INTO factures (type,titre,date,montant,fichier_path,sale_id)
                             VALUES (?,?,?,?,?,?)""",
                         ("client", f"{client_label} ({format_label})", today_iso(), s["vente"], path, sid))
        conn.commit()
        conn.close()
        log_action(self.app.current_user["username"], f"Facture {format_label} générée : {c['nom'] if c else '?'}")
        open_file(path)


def _draw_invoice_pdf(c, sale, client, rx):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    from reportlab.lib.units import mm
    from reportlab.lib.colors import HexColor

    RED_C = HexColor("#D6293B")
    BLACK_C = HexColor("#0B0B0C")

    
    w, h = A4

    def draw_logo(x, y, align="left"):
        """Petit pictogramme lunettes + mot-symbole OPTILUX, comme sur le papier à en-tête."""
        r = 3.6 * mm
        gap = 1.6 * mm
        if align == "right":
            cx2 = x
            cx1 = cx2 - 2 * r - gap
        else:
            cx1 = x
            cx2 = cx1 + 2 * r + gap
        c.setLineWidth(1)
        c.setStrokeColor(BLACK_C)
        c.circle(cx1, y, r, stroke=1, fill=0)
        c.setStrokeColor(RED_C)
        c.circle(cx2, y, r, stroke=1, fill=0)
        c.setStrokeColor(BLACK_C)
        c.line(cx1 + r * 0.4, y + r * 0.3, cx2 - r * 0.4, y + r * 0.3)
        text_x = (cx2 + r + 2 * mm) if align == "left" else (cx1 - r - 2 * mm)
        c.setFont("Helvetica-Bold", 13)
        if align == "left":
            c.setFillColor(BLACK_C)
            c.drawString(text_x, y - 4, "OPTI")
            optiw = c.stringWidth("OPTI", "Helvetica-Bold", 13)
            c.setFillColor(RED_C)
            c.drawString(text_x + optiw, y - 4, "LUX")
        else:
            c.setFillColor(RED_C)
            luxw = c.stringWidth("LUX", "Helvetica-Bold", 13)
            c.drawString(text_x - luxw, y - 4, "LUX")
            c.setFillColor(BLACK_C)
            optiw = c.stringWidth("OPTI", "Helvetica-Bold", 13)
            c.drawString(text_x - luxw - optiw, y - 4, "OPTI")
        c.setFillColor(BLACK_C)

    # ---------- en-tête ----------
    top_y = h - 18 * mm
    draw_logo(20 * mm, top_y, align="left")
    draw_logo(w - 20 * mm, top_y, align="right")

    year = (sale["date"] or today_iso())[:4]
    c.setFont("Helvetica-Bold", 14)
    c.setFillColor(RED_C)
    label = "Facture N° "
    label_w = c.stringWidth(label, "Helvetica-Bold", 14)
    num_str = f"{sale['id']}"
    num_w = c.stringWidth(num_str, "Helvetica-Bold", 14)
    suffix = f" / {year}"
    suffix_w = c.stringWidth(suffix, "Helvetica-Bold", 14)
    total_w = label_w + num_w + suffix_w
    start_x = w / 2 - total_w / 2
    c.drawString(start_x, top_y - 2, label)
    c.setFillColor(BLACK_C)
    c.drawString(start_x + label_w, top_y - 2, num_str)
    c.drawString(start_x + label_w + num_w, top_y - 2, suffix)

    c.setFont("Helvetica", 10)
    c.drawCentredString(w / 2, top_y - 16, "OPTICIENNE - OPTOMETRISTE")
    c.setFont("Helvetica-Bold", 12)
    c.drawCentredString(w / 2, top_y - 32, "HAOURIR Sanae")

    date_str = sale["date"] or today_iso()
    try:
        yyyy, mm_, dd = date_str.split("-")
    except ValueError:
        yyyy, mm_, dd = year, "", ""
    c.setFont("Helvetica", 10)
    c.drawCentredString(w / 2, top_y - 46, f"Tanger le : {dd} / {mm_} / {yyyy}")

    # ---------- client ----------
    y = top_y - 64
    nom_complet = f"{client['nom']} {client['prenom'] or ''}".strip() if client else ""
    c.setFont("Helvetica", 10)
    c.drawString(20 * mm, y, "Mr – Mme – Mlle : ")
    prefix_w = c.stringWidth("Mr – Mme – Mlle : ", "Helvetica", 10)
    c.line(20 * mm + prefix_w, y - 1, w - 20 * mm, y - 1)
    c.setFont("Helvetica-Bold", 10)
    c.drawString(20 * mm + prefix_w + 2, y + 1, nom_complet)

    # ---------- Vision de LOIN / ADD Vision de PRES ----------
    y -= 10 * mm
    box_h = 22 * mm
    box_w = (w - 40 * mm - 6 * mm) / 2
    box1_x = 20 * mm
    box2_x = box1_x + box_w + 6 * mm

    def vision_box(x, title, od_val, og_val):
        c.setLineWidth(1)
        c.setStrokeColor(BLACK_C)
        c.rect(x, y - box_h, box_w, box_h)
        c.line(x, y - 8 * mm, x + box_w, y - 8 * mm)
        c.setFont("Helvetica-Bold", 10)
        c.drawCentredString(x + box_w / 2, y - 5.5 * mm, title)
        c.setFont("Helvetica-Bold", 9)
        c.drawString(x + 3 * mm, y - 13.5 * mm, "OD :")
        c.drawString(x + 3 * mm, y - 19.5 * mm, "OG :")
        c.setFont("Helvetica", 9)
        c.drawString(x + 12 * mm, y - 13.5 * mm, od_val or "")
        c.drawString(x + 12 * mm, y - 19.5 * mm, og_val or "")
        c.line(x + 11 * mm, y - 14.3 * mm, x + box_w - 3 * mm, y - 14.3 * mm)
        c.line(x + 11 * mm, y - 20.3 * mm, x + box_w - 3 * mm, y - 20.3 * mm)

    def fmt_loin(sph, cyl, axe):
        if not sph and not cyl and not axe:
            return ""
        return f"{sph or '---'} ({cyl or '---'}) {axe or ''}".strip()

    loin_od = fmt_loin(rx["od_sph"], rx["od_cyl"], rx["od_axe"]) if rx else ""
    loin_og = fmt_loin(rx["og_sph"], rx["og_cyl"], rx["og_axe"]) if rx else ""
    pres_od = (rx["od_add"] or "") if rx else ""
    pres_og = (rx["og_add"] or "") if rx else ""

    vision_box(box1_x, "Vision de LOIN", loin_od, loin_og)
    vision_box(box2_x, "ADD Vision de PRES", pres_od, pres_og)

    # ---------- VL / VP / PROGRESSIF ----------
    y -= box_h + 10 * mm

    def checkbox(x, label, checked):
        size = 3.6 * mm
        c.setFont("Helvetica-Bold", 10)
        c.setFillColor(BLACK_C)
        c.drawString(x, y - 2.6, label)
        box_x = x + c.stringWidth(label, "Helvetica-Bold", 10) + 2
        c.setLineWidth(1)
        c.rect(box_x, y - 3, size, size)
        if checked:
            c.line(box_x, y - 3, box_x + size, y - 3 + size)
            c.line(box_x, y - 3 + size, box_x + size, y - 3)
        return box_x + size

    vl_checked = bool(rx["vl"]) if rx else False
    vp_checked = bool(rx["vp"]) if rx else False
    pg_checked = bool(rx["pg"]) if rx else False
    cx = w / 2 - 45 * mm
    cx = checkbox(cx, "VL", vl_checked) + 14 * mm
    cx = checkbox(cx, "VP", vp_checked) + 14 * mm
    checkbox(cx, "PROGRESSIF", pg_checked)

    # ---------- table Désignation / Prix unité / Montant TTC ----------
    y -= 12 * mm
    table_top = y
    col1_w = (w - 40 * mm) * 0.62
    col2_w = (w - 40 * mm) * 0.19
    col3_w = (w - 40 * mm) * 0.19
    x0 = 20 * mm
    x1 = x0 + col1_w
    x2 = x1 + col2_w
    x3 = x2 + col3_w
    header_h = 10 * mm
    rows_h = 24 * mm
    bottom_h = 20 * mm
    total_table_h = header_h + rows_h + bottom_h

    c.setLineWidth(1)
    c.rect(x0, table_top - total_table_h, x3 - x0, total_table_h)
    c.line(x0, table_top - header_h, x3, table_top - header_h)
    c.line(x0, table_top - header_h - rows_h, x3, table_top - header_h - rows_h)
    c.line(x1, table_top, x1, table_top - total_table_h)
    c.line(x2, table_top, x2, table_top - total_table_h)

    c.setFont("Helvetica-Bold", 9)
    c.drawCentredString((x0 + x1) / 2, table_top - 6.5 * mm, "Désignation")
    c.drawCentredString((x1 + x2) / 2, table_top - 6.5 * mm, "Prix unité")
    c.drawCentredString((x2 + x3) / 2, table_top - 6.5 * mm, "Montant TTC")

    verres_desc = (rx["verres_texte"] if rx and rx["verres_texte"] else "") or "correcteurs organiques antireflet"
    c.setFont("Helvetica", 9)
    c.drawString(x0 + 3 * mm, table_top - header_h - 7 * mm, "-   Monture optique.")
    c.drawString(x0 + 3 * mm, table_top - header_h - 15 * mm, f"-   Verres correcteurs : {verres_desc}")

    c.setFont("Helvetica", 9)
    c.drawCentredString((x1 + x2) / 2, table_top - header_h - 7 * mm, f"{sale['monture']:.0f},00")
    c.drawCentredString((x2 + x3) / 2, table_top - header_h - 7 * mm, f"{sale['monture']:.0f},00")
    c.drawCentredString((x1 + x2) / 2, table_top - header_h - 19 * mm, f"{sale['verres']:.0f},00")
    c.drawCentredString((x2 + x3) / 2, table_top - header_h - 19 * mm, f"{sale['verres']:.0f},00")

    bottom_y = table_top - header_h - rows_h
    c.setFont("Helvetica-Bold", 9)
    c.drawString(x0 + 3 * mm, bottom_y - 7 * mm, "Cachet et signature :")
    c.drawCentredString((x1 + x2) / 2, bottom_y - 7 * mm, "TOTAL")
    c.drawCentredString((x1 + x2) / 2, bottom_y - 13 * mm, "TTC")
    remise_txt = f" (-{sale['remise']}%)" if sale["remise"] else ""
    c.setFont("Helvetica-Bold", 10)
    c.drawCentredString((x2 + x3) / 2, bottom_y - 10 * mm, f"{sale['vente']:.0f},00")
    if remise_txt:
        c.setFont("Helvetica", 7)
        c.drawCentredString((x2 + x3) / 2, bottom_y - 16 * mm, remise_txt.strip())

    # ---------- montant en toutes lettres ----------
    y = table_top - total_table_h - 12 * mm
    words = num_to_french_words(sale["vente"]).upper()
    c.setFont("Helvetica-Bold", 10)
    c.drawString(20 * mm, y, "\u2713  Arrêtée la présente facture à la somme de : ")
    prefix_w = c.stringWidth("\u2713  Arrêtée la présente facture à la somme de : ", "Helvetica-Bold", 10)
    c.setFont("Helvetica", 10)
    c.drawString(20 * mm + prefix_w, y, f"{words} DIRHAMS.")
    c.line(20 * mm, y - 8, w - 20 * mm, y - 8)

    y -= 20
    c.setFont("Helvetica-Oblique", 8)
    c.drawCentredString(w / 2, y, "NB : Tous les montants sont exprimés en Dirhams.")

    # ---------- pied de page ----------
    y -= 16
    c.setLineWidth(1)
    c.rect(20 * mm, y - 18, w - 40 * mm, 18)
    c.setFont("Helvetica", 7)
    c.drawCentredString(w / 2, y - 8,
                         "ADRESSE : VIENNA MALL - B1 angle Moutanabi et Ahmed chawki - RDC N°04 - TANGER")
    c.drawCentredString(w / 2, y - 15,
                         "RC: 148206 / PATENTE: 50405209 / IF: 24813627 / ICE: 001962144000020 / NUM : 06 49 24 94 24")

def generate_facture_pdf(sale, client, rx):
    """A4 landscape PDF with two identical invoices side by side (customer + shop copy)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    from reportlab.lib.units import mm

    out_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "factures")
    os.makedirs(out_dir, exist_ok=True)
    client_label = f"{client['nom']}_{client['prenom'] or ''}".strip("_") if client else "client"
    fname = f"Facture_{sale['id']}_{client_label.replace(' ', '_')}.pdf"
    path = os.path.join(out_dir, fname)

    page_w, page_h = A4[1], A4[0]
    logical_w, logical_h = A4
    half_w = page_w / 2.0

    scale = min(half_w / logical_w, page_h / logical_h) * 0.96
    scaled_w = logical_w * scale
    scaled_h = logical_h * scale
    x_offset = (half_w - scaled_w) / 2.0
    y_offset = (page_h - scaled_h) / 2.0

    c = canvas.Canvas(path, pagesize=(page_w, page_h))

    c.saveState()
    c.translate(x_offset, y_offset)
    c.scale(scale, scale)
    _draw_invoice_pdf(c, sale, client, rx)
    c.restoreState()

    c.saveState()
    c.translate(half_w + x_offset, y_offset)
    c.scale(scale, scale)
    _draw_invoice_pdf(c, sale, client, rx)
    c.restoreState()

    c.setDash(3, 3)
    c.setStrokeColorRGB(0.65, 0.65, 0.65)
    c.setLineWidth(0.4)
    c.line(half_w, 6 * mm, half_w, page_h - 6 * mm)
    c.setDash()

    c.save()
    return path

def _draw_invoice_docx(container, sale, client, rx):
    from docx import Document
    from docx.shared import Pt, Mm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    RED_RGB = RGBColor(0xD6, 0x29, 0x3B)
    BLACK_RGB = RGBColor(0x0B, 0x0B, 0x0C)

    def set_cell_border(cell, **kwargs):
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_borders = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right"):
            if edge in kwargs:
                el = OxmlElement(f"w:{edge}")
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), "6")
                el.set(qn("w:color"), "0B0B0C")
                tc_borders.append(el)
        tc_pr.append(tc_borders)

    # ---------- en-tête (logo texte + titre) ----------
    header_table = container.add_table(rows=1, cols=3)
    header_table.autofit = True
    logo_path = os.path.join(ASSETS_DIR, "optilux_logo.png")
    for col_idx, align in ((0, WD_ALIGN_PARAGRAPH.LEFT), (2, WD_ALIGN_PARAGRAPH.RIGHT)):
        cell = header_table.rows[0].cells[col_idx]
        p = cell.paragraphs[0]
        p.alignment = align
        if os.path.exists(logo_path):
            try:
                p.add_run().add_picture(logo_path, height=Mm(9))
            except Exception:
                pass
        run = p.add_run(" OPTILUX")
        run.bold = True
        run.font.size = Pt(13)

    year = (sale["date"] or today_iso())[:4]
    p = container.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"Facture N° {sale['id']} / {year}")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RED_RGB

    for text, size, bold in (("OPTICIENNE - OPTOMETRISTE", 11, False), ("HAOURIR Sanae", 13, True)):
        p = container.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)

    date_str = sale["date"] or today_iso()
    try:
        yyyy, mm_, dd = date_str.split("-")
    except ValueError:
        yyyy, mm_, dd = year, "", ""
    p = container.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"Tanger le : {dd} / {mm_} / {yyyy}")

    container.add_paragraph()
    nom_complet = f"{client['nom']} {client['prenom'] or ''}".strip() if client else ""
    p = container.add_paragraph()
    run = p.add_run("Mr – Mme – Mlle : ")
    run2 = p.add_run(nom_complet)
    run2.bold = True
    run2.underline = True

    # ---------- Vision de LOIN / ADD Vision de PRES ----------
    def fmt_loin(sph, cyl, axe):
        if not sph and not cyl and not axe:
            return ""
        return f"{sph or '---'} ({cyl or '---'}) {axe or ''}".strip()

    loin_od = fmt_loin(rx["od_sph"], rx["od_cyl"], rx["od_axe"]) if rx else ""
    loin_og = fmt_loin(rx["og_sph"], rx["og_cyl"], rx["og_axe"]) if rx else ""
    pres_od = (rx["od_add"] or "") if rx else ""
    pres_og = (rx["og_add"] or "") if rx else ""

    vt = container.add_table(rows=3, cols=2)
    vt.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = vt.rows[0].cells
    headers[0].text = "Vision de LOIN"
    headers[1].text = "ADD Vision de PRES"
    for cell in headers:
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_border(cell, top=1, left=1, right=1, bottom=1)
    for row_i, (od_val, og_val_pair) in enumerate([
        (loin_od, pres_od), (loin_og, pres_og)
    ], start=1):
        c0 = vt.rows[row_i].cells[0]
        c1 = vt.rows[row_i].cells[1]
        label = "OD : " if row_i == 1 else "OG : "
        c0.text = f"{label}{od_val}"
        c1.text = f"{label}{og_val_pair}"
        for cell in (c0, c1):
            set_cell_border(cell, top=1, left=1, right=1, bottom=1)

    container.add_paragraph()
    vl_checked = bool(rx["vl"]) if rx else False
    vp_checked = bool(rx["vp"]) if rx else False
    pg_checked = bool(rx["pg"]) if rx else False
    p = container.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for label, checked in (("VL", vl_checked), ("VP", vp_checked), ("PROGRESSIF", pg_checked)):
        box = "\u2611" if checked else "\u2610"  # ☑ / ☐ — glyphes Unicode standards, rendus nativement par Word
        run = p.add_run(f"  {label} {box}   ")
        run.bold = True

    # ---------- table Désignation / Prix unité / Montant TTC ----------
    container.add_paragraph()
    dt = container.add_table(rows=3, cols=3)
    dt.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = dt.rows[0].cells
    for i, txt in enumerate(("Désignation", "Prix unité", "Montant TTC")):
        hdr[i].text = txt
        hdr[i].paragraphs[0].runs[0].bold = True
        hdr[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    verres_desc = (rx["verres_texte"] if rx and rx["verres_texte"] else "") or "correcteurs organiques antireflet"
    row1 = dt.rows[1].cells
    row1[0].text = "-  Monture optique."
    row1[1].text = f"{sale['monture']:.0f},00"
    row1[2].text = f"{sale['monture']:.0f},00"

    row2 = dt.rows[2].cells
    row2[0].text = f"-  Verres correcteurs : {verres_desc}"
    row2[1].text = f"{sale['verres']:.0f},00"
    row2[2].text = f"{sale['verres']:.0f},00"

    for row in dt.rows:
        for cell in row.cells:
            set_cell_border(cell, top=1, left=1, right=1, bottom=1)
            for para in cell.paragraphs:
                if not para.alignment and cell != row1[0] and cell != row2[0]:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    tt = container.add_table(rows=1, cols=2)
    tt.alignment = WD_TABLE_ALIGNMENT.CENTER
    c0, c1 = tt.rows[0].cells
    c0.text = "Cachet et signature :"
    remise_txt = f"  (-{sale['remise']}%)" if sale["remise"] else ""
    c1.text = f"TOTAL TTC : {sale['vente']:.0f},00 DH{remise_txt}"
    c1.paragraphs[0].runs[0].bold = True
    for cell in (c0, c1):
        set_cell_border(cell, top=1, left=1, right=1, bottom=1)

    # ---------- montant en lettres ----------
    container.add_paragraph()
    words = num_to_french_words(sale["vente"]).upper()
    p = container.add_paragraph()
    run = p.add_run("\u2713  Arrêtée la présente facture à la somme de : ")
    run.bold = True
    p.add_run(f"{words} DIRHAMS.")

    p = container.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("NB : Tous les montants sont exprimés en Dirhams.")
    run.italic = True
    run.font.size = Pt(9)

    container.add_paragraph()
    p = container.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("ADRESSE : VIENNA MALL - B1 angle Moutanabi et Ahmed chawki - RDC N°04 - TANGER")
    run.font.size = Pt(8)
    p2 = container.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run("RC: 148206 / PATENTE: 50405209 / IF: 24813627 / ICE: 001962144000020 / NUM : 06 49 24 94 24")
    run2.font.size = Pt(8)

TEMPLATE_DOCX = os.path.join(os.path.dirname(os.path.abspath(__file__)), "modell2027.docx")

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _q(tag):
    return f"{{{W_NS}}}{tag}"


def _fill_content_controls(template_path, output_path, values):
    """Fill Word content controls (w:sdt) by their <w:tag> value.
    `values` maps tag name → replacement text. Names not in the map stay untouched.
    Content controls whose tag appears twice (two side-by-side copies) get the
    same value in both places."""
    shutil.copy(template_path, output_path)

    with zipfile.ZipFile(template_path, "r") as z:
        doc_xml_bytes = z.read("word/document.xml")

    root = etree.fromstring(doc_xml_bytes)
    xml_space = "{http://www.w3.org/XML/1998/namespace}space"
    filled = set()

    for sdt in root.iter(_q("sdt")):
        sdtPr = sdt.find(_q("sdtPr"))
        if sdtPr is None:
            continue
        tag_el = sdtPr.find(_q("tag"))
        if tag_el is None:
            continue
        name = tag_el.get(_q("val"))
        if name not in values:
            continue
        content = sdt.find(_q("sdtContent"))
        if content is None:
            continue
        t_els = list(content.iter(_q("t")))
        if not t_els:
            continue
        t_els[0].text = str(values[name])
        t_els[0].set(xml_space, "preserve")
        for t in t_els[1:]:
            t.text = ""
        filled.add(name)

    new_xml = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    tmp_out = output_path + ".tmp"
    with zipfile.ZipFile(tmp_out, "w", zipfile.ZIP_DEFLATED) as zout:
        with zipfile.ZipFile(template_path, "r") as zin:
            for item in zin.namelist():
                data = new_xml if item == "word/document.xml" else zin.read(item)
                zout.writestr(item, data)
    shutil.move(tmp_out, output_path)
    return filled


def generate_facture_docx(sale, client, rx):
    """Fills the Word template's content controls with the sale's data
    and saves the result under /factures/."""
    if not os.path.exists(TEMPLATE_DOCX):
        raise FileNotFoundError(
            f"Modèle Word introuvable : {TEMPLATE_DOCX}\n"
            "Placez modell2027.docx à côté de app.py."
        )

    out_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "factures")
    os.makedirs(out_dir, exist_ok=True)
    client_label = f"{client['nom']}_{client['prenom'] or ''}".strip("_") if client else "client"
    fname = f"Facture_{sale['id']}_{client_label.replace(' ', '_')}.docx"
    path = os.path.join(out_dir, fname)

    # ---------- prep values ----------
    date_str = sale["date"] or today_iso()
    try:
        yyyy, mm_, dd = date_str.split("-")
    except ValueError:
        yyyy, mm_, dd = date_str[:4], "", ""

    def fmt_loin(sph, cyl, axe):
        if not sph and not cyl and not axe:
            return ""
        return f"{sph or '---'} ({cyl or '---'}) {axe or ''}".strip()

    loin_od = fmt_loin(rx["od_sph"], rx["od_cyl"], rx["od_axe"]) if rx else ""
    loin_og = fmt_loin(rx["og_sph"], rx["og_cyl"], rx["og_axe"]) if rx else ""
    pres_od = (rx["od_add"] or "") if rx else ""
    pres_og = (rx["og_add"] or "") if rx else ""

    # ---------- mapping (names must match the content control tags in the template) ----------
    values = {
        "numero":     str(sale["id"]),
        "date_jour":  date_str,
        "client_1":   client["nom"] if client else "client",
        "loin_od_1":  loin_od,
        "loin_og_1":  loin_og,
        "pres_od_1":  pres_od,
        "pres_og_1":  pres_og,
        "vl_1":         "☑" if bool(rx["vl"]) else "☐",
        "vp_1":         "☑" if bool(rx["vp"]) else "☐",
        "progressif_1":   "☑" if bool(rx["pg"]) else "☐",
        "prix_monture_1": f"{sale['monture']:.0f},00",
        "montant_monture_1": f"{sale['monture']:.0f},00",
        "prix_verres_1":  f"{sale['verres']:.0f},00",
        "montant_verres_1": f"{sale['verres']:.0f},00",
        "total_1":        f"{sale['vente']:.0f},00",
        "somme_lettres_1": num_to_french_words(sale["vente"]).upper(),
        
    }

    # ---------- render ----------
    filled = _fill_content_controls(TEMPLATE_DOCX, path, values)

    # Optional: warn on the console if a field wasn't found.
    missing = set(values) - filled
    if missing:
        print(f"[Word] Champs non trouvés dans le modèle : {', '.join(sorted(missing))}")

    return path


# ================================================================== EXPENSES (admin only)
class ExpensesPage(BasePage):
    def __init__(self, parent, app):
        super().__init__(parent, app)
        self.show_personal = False
        self.header("Charges", "Dépenses", "+ Nouvelle dépense", self.open_form)

        toggle_row = tk.Frame(self, bg=CONTENT_BG)
        toggle_row.pack(fill="x", pady=(0, 10))
        self.btn_business = self.mode_button(toggle_row, "Charges (entreprise)", lambda: self.set_mode(False))
        self.btn_business.pack(side="left", padx=(0, 8))
        self.btn_personal = self.mode_button(toggle_row, "Dépenses personnelles", lambda: self.set_mode(True))
        self.btn_personal.pack(side="left")

        self.search_var = self.make_search_box("Rechercher une dépense…")
        self.tree = self.make_table(["Date", "Libellé", "Catégorie", "Montant", "Échéance", "Récurrent", "Payé"])
        self.tree.tag_configure("paid", background="#E8F5EC", foreground="#2E7D46")
        self.tree.bind("<Double-1>", lambda e: self.open_form(self._selected_id()))
        btns = tk.Frame(self, bg=CONTENT_BG)
        btns.pack(fill="x", pady=(8, 0))
        outline_button(btns, "Modifier la sélection", lambda: self.open_form(self._selected_id()), icon="edit").pack(side="left")
        danger_button(btns, "Supprimer", self.delete_selected).pack(side="left", padx=(0, 12))
        self._update_toggle_style()

    def set_mode(self, personal):
        self.show_personal = personal
        self._update_toggle_style()
        self.refresh()

    def _update_toggle_style(self):
        self.set_mode_button_active(self.btn_personal, self.show_personal)
        self.set_mode_button_active(self.btn_business, not self.show_personal)

    def _selected_id(self):
        sel = self.tree.selection()
        return int(sel[0]) if sel else None

    ALERT_DAYS = 3

    def refresh(self):
        for i in self.tree.get_children():
            self.tree.delete(i)
        q = self.search_var.get()
        conn = get_connection()
        rows = conn.execute("SELECT * FROM expenses WHERE personnelle=? ORDER BY date DESC",
                             (int(self.show_personal),)).fetchall()
        conn.close()
        today = datetime.date.today()
        n = 0
        for x in rows:
            if not self.row_matches(q, x["libelle"], x["categorie"]):
                continue
            paid = expense_is_paid_this_month(x)
            if paid:
                tag = "paid"
            elif x["date_echeance"]:
                try:
                    days_left = (datetime.date.fromisoformat(x["date_echeance"]) - today).days
                except ValueError:
                    days_left = None
                if days_left is not None and days_left <= self.ALERT_DAYS:
                    tag = "low"
                else:
                    tag = "even" if n % 2 == 0 else "odd"
            else:
                tag = "even" if n % 2 == 0 else "odd"
            self.tree.insert("", "end", iid=str(x["id"]), tags=(tag,),
                              values=(x["date"], x["libelle"], x["categorie"], money(x["montant"]),
                                      x["date_echeance"] or "—", "Oui" if x["recurrent"] else "Non",
                                      "Oui" if paid else "Non"))
            n += 1

    def delete_selected(self):
        xid = self._selected_id()
        if not xid:
            messagebox.showinfo("Info", "Sélectionnez une dépense.")
            return
        if not messagebox.askyesno("Confirmer", "Supprimer cette dépense ?"):
            return
        conn = get_connection()
        row = conn.execute("SELECT libelle FROM expenses WHERE id=?", (xid,)).fetchone()
        conn.execute("DELETE FROM expenses WHERE id=?", (xid,))
        conn.commit(); conn.close()
        log_action(self.app.current_user["username"], f"Dépense supprimée : {row['libelle']}")
        self.refresh()

    def open_form(self, exp_id=None):
        conn = get_connection()
        x = conn.execute("SELECT * FROM expenses WHERE id=?", (exp_id,)).fetchone() if exp_id else None
        conn.close()
        cats = get_types("depense_categorie") or ["Autre"]

        modal = ModalForm(self.app, "Modifier dépense" if x else "Nouvelle dépense", width=400)
        body = modal.body
        libelle_e = field(body, "Libellé", x["libelle"] if x else "")
        cat_v = dropdown(body, "Catégorie", cats, x["categorie"] if x else cats[0])
        montant_e = field(body, "Montant (DH)", str(x["montant"]) if x else "0")
        date_e = field_date(body, "Date", x["date"] if x else today_iso())
        echeance_e = field_date(body, "Échéance (rappel, optionnel)", x["date_echeance"] if x and x["date_echeance"] else "")
        recur_var = tk.BooleanVar(value=bool(x["recurrent"]) if x else False)
        tk.Checkbutton(body, text="Dépense récurrente (mensuelle)", variable=recur_var, bg=SURFACE,
                       fg=TEXT, activebackground=SURFACE, activeforeground=TEXT,
                       selectcolor=SURFACE, font=FONT_BODY).pack(anchor="w", pady=(10, 0))
        personal_var = tk.BooleanVar(value=bool(x["personnelle"]) if x else self.show_personal)
        tk.Checkbutton(body, text="Dépense personnelle (hors entreprise)", variable=personal_var, bg=SURFACE,
                       fg=TEXT, activebackground=SURFACE, activeforeground=TEXT,
                       selectcolor=SURFACE, font=FONT_BODY).pack(anchor="w", pady=(4, 0))
        paid_var = oui_non(body, "Payé",
                        expense_is_paid_this_month(x) if x else False)

        def save():
            libelle = libelle_e.get().strip()
            if not libelle:
                messagebox.showwarning("Champ requis", "Le libellé est obligatoire.")
                return
            try:
                montant = float(montant_e.get() or 0)
            except ValueError:
                messagebox.showwarning("Valeur invalide", "Montant invalide.")
                return
            conn = get_connection()
            paid_val = 1 if paid_var.get() == "oui" else 0
            # For a recurring charge marked paid, remember *which* month this payment
            # covers. For non-recurring or unpaid, leave paye_mois as it was / null.
            if paid_val and int(recur_var.get()):
                paye_mois_val = datetime.date.today().strftime("%Y-%m")
            elif paid_val:
                paye_mois_val = None  # one-off; paye_mois is unused
            else:
                # marked unpaid → clear the month marker too
                paye_mois_val = None

            if x:
                conn.execute("""UPDATE expenses SET libelle=?, categorie=?, montant=?, date=?, recurrent=?,
                                 date_echeance=?, personnelle=?, paye=?, paye_mois=? WHERE id=?""",
                             (libelle, cat_v.get(), montant, date_e.get(), int(recur_var.get()),
                              echeance_e.get() or None, int(personal_var.get()), paid_val,
                              paye_mois_val, x["id"]))
                action = f"Dépense modifiée : {libelle}"
            else:
                conn.execute("""INSERT INTO expenses (date,libelle,categorie,montant,recurrent,
                                 date_echeance,personnelle,paye,paye_mois)
                                 VALUES (?,?,?,?,?,?,?,?,?)""",
                             (date_e.get(), libelle, cat_v.get(), montant, int(recur_var.get()),
                              echeance_e.get() or None, int(personal_var.get()), paid_val,
                              paye_mois_val))
                action = f"Dépense créée : {libelle} ({money(montant)})"

            conn.commit(); conn.close()
            log_action(self.app.current_user["username"], action)
            modal.destroy()
            self.set_mode(bool(personal_var.get()))

        modal.buttons(save)


# ================================================================== USERS & AUDIT (admin only)
class UsersPage(BasePage):
    def __init__(self, parent, app):
        super().__init__(parent, app)
        self.header("Accès", "Utilisateurs", "+ Nouvel utilisateur", self.open_form)
        self.tree = self.make_table(["Identifiant", "Nom", "Rôle"], height=6)
        self.tree.bind("<Double-1>", lambda e: self.open_form(self._selected_id()))
        btns = tk.Frame(self, bg=CONTENT_BG)
        btns.pack(fill="x", pady=(8, 14))
        outline_button(btns, "Modifier la sélection", lambda: self.open_form(self._selected_id()), icon="edit").pack(side="left")
        danger_button(btns, "Supprimer", self.delete_selected).pack(side="left", padx=(0, 12))

        tk.Label(self, text="Journal d'activité", font=FONT_DISPLAY_SM, bg=CONTENT_BG, fg=TEXT).pack(anchor="w", pady=(6, 8))
        self.audit_search_var = self.make_search_box("Rechercher dans le journal (utilisateur, action)…")
        self.audit_tree = self.make_table(["Date / heure", "Utilisateur", "Action"], height=10)

    def _selected_id(self):
        sel = self.tree.selection()
        return int(sel[0]) if sel else None

    def refresh(self):
        for i in self.tree.get_children():
            self.tree.delete(i)
        conn = get_connection()
        rows = conn.execute("SELECT * FROM users ORDER BY username").fetchall()
        for n, u in enumerate(rows):
            self.tree.insert("", "end", iid=str(u["id"]), tags=("even" if n % 2 == 0 else "odd",),
                              values=(u["username"], u["nom"], "Administrateur" if u["role"] == "admin" else "Employé"))
        for i in self.audit_tree.get_children():
            self.audit_tree.delete(i)
        audit = conn.execute("SELECT * FROM audit_log ORDER BY id DESC LIMIT 200").fetchall()
        conn.close()
        q = self.audit_search_var.get()
        n = 0
        for a in audit:
            if not self.row_matches(q, a["user"], a["action"]):
                continue
            self.audit_tree.insert("", "end", tags=("even" if n % 2 == 0 else "odd",),
                                    values=(a["ts"].replace("T", " "), a["user"], a["action"]))
            n += 1

    def delete_selected(self):
        uid = self._selected_id()
        if not uid:
            messagebox.showinfo("Info", "Sélectionnez un utilisateur.")
            return
        if uid == self.app.current_user["id"]:
            messagebox.showwarning("Action impossible", "Vous ne pouvez pas supprimer votre propre compte.")
            return
        conn = get_connection()
        target = conn.execute("SELECT * FROM users WHERE id=?", (uid,)).fetchone()
        admins = conn.execute("SELECT COUNT(*) as n FROM users WHERE role='admin'").fetchone()["n"]
        if target["role"] == "admin" and admins <= 1:
            messagebox.showwarning("Action impossible", "Impossible de supprimer le dernier compte administrateur.")
            conn.close()
            return
        if not messagebox.askyesno("Confirmer", "Supprimer cet utilisateur ?"):
            conn.close()
            return
        conn.execute("DELETE FROM users WHERE id=?", (uid,))
        conn.commit(); conn.close()
        log_action(self.app.current_user["username"], f"Utilisateur supprimé : {target['username']}")
        self.refresh()

    def open_form(self, user_id=None):
        conn = get_connection()
        u = conn.execute("SELECT * FROM users WHERE id=?", (user_id,)).fetchone() if user_id else None
        conn.close()

        modal = ModalForm(self.app, "Modifier utilisateur" if u else "Nouvel utilisateur", width=380)
        body = modal.body
        nom_e = field(body, "Nom complet", u["nom"] if u else "")
        username_e = field(body, "Identifiant", u["username"] if u else "")
        pass_e = field(body, "Mot de passe" + (" (laisser vide pour ne pas changer)" if u else ""), "", show="•")
        role_v = dropdown(body, "Rôle", ["employee", "admin"], u["role"] if u else "employee")

        def save():
            nom = nom_e.get().strip()
            username = username_e.get().strip()
            if not nom or not username:
                messagebox.showwarning("Champs requis", "Nom et identifiant sont obligatoires.")
                return
            from auth import hash_password
            conn = get_connection()
            existing = conn.execute("SELECT id FROM users WHERE username=? AND id!=?",
                                     (username, u["id"] if u else -1)).fetchone()
            if existing:
                messagebox.showwarning("Identifiant pris", "Cet identifiant existe déjà.")
                conn.close()
                return
            if u:
                if pass_e.get():
                    h, salt = hash_password(pass_e.get())
                    conn.execute("UPDATE users SET nom=?, username=?, password_hash=?, salt=?, role=? WHERE id=?",
                                 (nom, username, h, salt, role_v.get(), u["id"]))
                else:
                    conn.execute("UPDATE users SET nom=?, username=?, role=? WHERE id=?",
                                 (nom, username, role_v.get(), u["id"]))
                action = f"Utilisateur modifié : {username}"
            else:
                if not pass_e.get():
                    messagebox.showwarning("Champ requis", "Le mot de passe est obligatoire.")
                    conn.close()
                    return
                h, salt = hash_password(pass_e.get())
                conn.execute("INSERT INTO users (username,password_hash,salt,nom,role) VALUES (?,?,?,?,?)",
                             (username, h, salt, nom, role_v.get()))
                action = f"Utilisateur créé : {username}"
            conn.commit(); conn.close()
            log_action(self.app.current_user["username"], action)
            modal.destroy()
            self.refresh()
            messagebox.showinfo("Redémarrage requis",
                                 "La liste des menus (admin/employé) se met à jour à la prochaine connexion.")

        modal.buttons(save)

# ================================================================== TYPES & PRODUITS (admin only)
# ================================================================== FACTURES
class FacturesPage(BasePage):
    def __init__(self, parent, app):
        super().__init__(parent, app)
        self.show_type = "client"
        self.header("Documents", "Factures", "+ Ajouter manuellement", self.open_form)

        toggle_row = tk.Frame(self, bg=CONTENT_BG)
        toggle_row.pack(fill="x", pady=(0, 10))
        self.btn_client = self.mode_button(toggle_row, "Factures clients", lambda: self.set_type("client"))
        self.btn_client.pack(side="left", padx=(0, 8))
        self.btn_entreprise = self.mode_button(toggle_row, "Factures de l'entreprise",
                                                lambda: self.set_type("entreprise"))
        self.btn_entreprise.pack(side="left")

        self.search_var = self.make_search_box("Rechercher une facture…")
        self.tree = self.make_table(["Titre", "Date", "Montant", "Fichier"])
        self.tree.bind("<Double-1>", lambda e: self.open_selected_file())
        btns = tk.Frame(self, bg=CONTENT_BG)
        btns.pack(fill="x", pady=(8, 0))
        outline_button(btns, "Ouvrir le fichier", self.open_selected_file, icon="document").pack(side="left")
        outline_button(btns, "Modifier", lambda: self.open_form(self._selected_id()), icon="edit").pack(side="left", padx=(12, 0))
        danger_button(btns, "Supprimer", self.delete_selected).pack(side="left", padx=(12, 0))
        self._update_toggle_style()

    def set_type(self, t):
        self.show_type = t
        self._update_toggle_style()
        self.refresh()

    def _update_toggle_style(self):
        self.set_mode_button_active(self.btn_client, self.show_type == "client")
        self.set_mode_button_active(self.btn_entreprise, self.show_type == "entreprise")

    def _selected_id(self):
        sel = self.tree.selection()
        return int(sel[0]) if sel else None

    def refresh(self):
        for i in self.tree.get_children():
            self.tree.delete(i)
        q = self.search_var.get()
        conn = get_connection()
        rows = conn.execute("SELECT * FROM factures WHERE type=? ORDER BY date DESC",
                             (self.show_type,)).fetchall()
        conn.close()
        n = 0
        for f in rows:
            if not self.row_matches(q, f["titre"], f["notes"]):
                continue
            self.tree.insert("", "end", iid=str(f["id"]), tags=("even" if n % 2 == 0 else "odd",),
                              values=(f["titre"] or "—", f["date"] or "—", money(f["montant"]),
                                      os.path.basename(f["fichier_path"]) if f["fichier_path"] else "—"))
            n += 1

    def open_selected_file(self):
        fid = self._selected_id()
        if not fid:
            messagebox.showinfo("Info", "Sélectionnez une facture.")
            return
        conn = get_connection()
        row = conn.execute("SELECT * FROM factures WHERE id=?", (fid,)).fetchone()
        conn.close()
        if not row or not row["fichier_path"] or not os.path.exists(row["fichier_path"]):
            messagebox.showinfo("Info", "Aucun fichier associé à cette facture.")
            return
        open_file(row["fichier_path"])

    def delete_selected(self):
        fid = self._selected_id()
        if not fid:
            messagebox.showinfo("Info", "Sélectionnez une facture.")
            return
        if not messagebox.askyesno("Confirmer", "Supprimer cette facture de la liste ?\n\n"
                                    "(Le fichier associé n'est pas supprimé du disque.)"):
            return
        conn = get_connection()
        conn.execute("DELETE FROM factures WHERE id=?", (fid,))
        conn.commit(); conn.close()
        log_action(self.app.current_user["username"], "Facture supprimée de la liste")
        self.refresh()

    def open_form(self, facture_id=None):
        conn = get_connection()
        f = conn.execute("SELECT * FROM factures WHERE id=?", (facture_id,)).fetchone() if facture_id else None
        conn.close()

        modal = ModalForm(self.app, "Modifier la facture" if f else "Ajouter une facture", width=420)
        body = modal.body
        titre_e = field(body, "Titre / Description", f["titre"] if f else "")
        row = tk.Frame(body, bg=SURFACE); row.pack(fill="x")
        row.grid_columnconfigure(0, weight=1); row.grid_columnconfigure(1, weight=1)
        date_wrap = tk.Frame(row, bg=SURFACE); date_wrap.grid(row=0, column=0, sticky="ew", padx=(0, 6))
        montant_wrap = tk.Frame(row, bg=SURFACE); montant_wrap.grid(row=0, column=1, sticky="ew", padx=(6, 0))
        date_e = field_date(date_wrap, "Date", f["date"] if f else today_iso())
        montant_e = field(montant_wrap, "Montant (DH)", str(f["montant"]) if f else "0")
        notes_e = field(body, "Notes (optionnel)", f["notes"] if f and f["notes"] else "")

        tk.Label(body, text="FICHIER", font=FONT_MONO_SM, fg=SLATE, bg=SURFACE, anchor="w").pack(fill="x", pady=(10, 2))
        state = {"path": f["fichier_path"] if f else None}
        file_lbl = tk.Label(body, text=os.path.basename(state["path"]) if state["path"] else "Aucun fichier choisi",
                             font=FONT_BODY, bg=SURFACE, fg=TEXT if state["path"] else SLATE, anchor="w")
        file_lbl.pack(fill="x", pady=(0, 6))

        def choose_file():
            path = filedialog.askopenfilename(
                title="Choisir un fichier (PDF ou image)",
                filetypes=[("Documents", "*.pdf *.png *.jpg *.jpeg *.webp"), ("Tous les fichiers", "*.*")])
            if not path:
                return
            try:
                os.makedirs(FACTURE_FILES_DIR, exist_ok=True)
                ext = os.path.splitext(path)[1]
                fname = f"{uuid.uuid4().hex}{ext}"
                dest = os.path.join(FACTURE_FILES_DIR, fname)
                with open(path, "rb") as src, open(dest, "wb") as dst:
                    dst.write(src.read())
                state["path"] = dest
                file_lbl.configure(text=os.path.basename(dest), fg=TEXT)
            except Exception as e:
                messagebox.showerror("Erreur", f"Impossible de copier le fichier :\n{e}")

        outline_button(body, "Choisir un fichier…", choose_file, icon="document").pack(fill="x")

        def save():
            titre = titre_e.get().strip()
            if not titre:
                messagebox.showwarning("Champ requis", "Le titre est obligatoire.")
                return
            try:
                montant = float(montant_e.get() or 0)
            except ValueError:
                messagebox.showwarning("Valeur invalide", "Montant invalide.")
                return
            conn = get_connection()
            if f:
                conn.execute("""UPDATE factures SET titre=?, date=?, montant=?, notes=?, fichier_path=?, type=?
                                 WHERE id=?""",
                             (titre, date_e.get(), montant, notes_e.get(), state["path"], self.show_type, f["id"]))
                action = f"Facture modifiée : {titre}"
            else:
                conn.execute("""INSERT INTO factures (type,titre,date,montant,fichier_path,notes)
                                 VALUES (?,?,?,?,?,?)""",
                             (self.show_type, titre, date_e.get(), montant, state["path"], notes_e.get()))
                action = f"Facture ajoutée : {titre}"
            conn.commit(); conn.close()
            log_action(self.app.current_user["username"], action)
            modal.destroy()
            self.refresh()

        modal.buttons(save)


class TypesPage(BasePage):
    def __init__(self, parent, app):
        super().__init__(parent, app)
        self.header("Paramètres", "Types & Produits")
        tk.Label(self, text="Gérez ici les listes utilisées dans les menus déroulants ailleurs dans "
                             "l'application (catégories, marques, types de lentilles…).",
                  font=FONT_BODY, bg=CONTENT_BG, fg=SLATE, wraplength=760, justify="left").pack(anchor="w", pady=(0, 16))

        picker_row = tk.Frame(self, bg=CONTENT_BG)
        picker_row.pack(fill="x", pady=(0, 12))
        tk.Label(picker_row, text="LISTE", font=FONT_MONO_SM, fg=SLATE, bg=CONTENT_BG).pack(side="left", padx=(0, 8))
        labels = [lbl for _key, lbl in TYPE_CATEGORIES]
        self.label_to_key = {lbl: key for key, lbl in TYPE_CATEGORIES}
        self.picker_var = tk.StringVar(value=labels[0])
        picker = ttk.Combobox(picker_row, textvariable=self.picker_var, values=labels, state="readonly",
                               font=FONT_BODY, style="Themed.TCombobox", width=32)
        picker.pack(side="left")
        picker.bind("<<ComboboxSelected>>", lambda e: self.refresh())

        add_row = tk.Frame(self, bg=CONTENT_BG)
        add_row.pack(fill="x", pady=(0, 12))
        self.new_value_entry = tk.Entry(add_row, font=FONT_BODY, relief="solid", borderwidth=1,
                                         bg=SURFACE, fg=TEXT, insertbackground=TEXT)
        self.new_value_entry.pack(side="left", fill="x", expand=True, ipady=6, padx=(0, 8))
        self.new_value_entry.bind("<Return>", lambda e: self.add_value())
        primary_button(add_row, "Ajouter", self.add_value, icon="plus").pack(side="left")

        self.tree = self.make_table(["Valeur"], height=12)
        btns = tk.Frame(self, bg=CONTENT_BG)
        btns.pack(fill="x", pady=(8, 0))
        danger_button(btns, "Supprimer la sélection", self.delete_selected).pack(side="left")

    def _current_categorie(self):
        return self.label_to_key.get(self.picker_var.get(), TYPE_CATEGORIES[0][0])

    def refresh(self):
        for i in self.tree.get_children():
            self.tree.delete(i)
        conn = get_connection()
        rows = conn.execute("SELECT * FROM types WHERE categorie=? ORDER BY valeur",
                             (self._current_categorie(),)).fetchall()
        conn.close()
        for n, t in enumerate(rows):
            self.tree.insert("", "end", iid=str(t["id"]), tags=("even" if n % 2 == 0 else "odd",),
                              values=(t["valeur"],))

    def add_value(self):
        val = self.new_value_entry.get().strip()
        if not val:
            return
        add_type(self._current_categorie(), val)
        log_action(self.app.current_user["username"],
                   f"Type ajouté : {val} ({self.picker_var.get()})")
        self.new_value_entry.delete(0, "end")
        self.refresh()

    def delete_selected(self):
        sel = self.tree.selection()
        if not sel:
            messagebox.showinfo("Info", "Sélectionnez une valeur à supprimer.")
            return
        type_id = int(sel[0])
        conn = get_connection()
        row = conn.execute("SELECT valeur FROM types WHERE id=?", (type_id,)).fetchone()
        conn.close()
        if not messagebox.askyesno("Confirmer", f"Supprimer « {row['valeur']} » ?\n\n"
                                    "Les articles existants qui utilisent déjà cette valeur ne seront pas modifiés."):
            return
        delete_type(type_id)
        log_action(self.app.current_user["username"], f"Type supprimé : {row['valeur']}")
        self.refresh()


# ================================================================== COMMANDES
# ================================================================== COMMANDES
class CommandesPage(BasePage):
    ALERT_DAYS = 2  # une commande est "proche" si sa date prévue est à J+ALERT_DAYS ou avant

    def __init__(self, parent, app):
        super().__init__(parent, app)
        self.show_type = "client"
        self.header("Suivi", "Commandes", "+ Nouvelle commande", self.open_form)

        toggle_row = tk.Frame(self, bg=CONTENT_BG)
        toggle_row.pack(fill="x", pady=(0, 10))
        self.btn_client = self.mode_button(toggle_row, "Commandes clients", lambda: self.set_type("client"))
        self.btn_client.pack(side="left", padx=(0, 8))
        self.btn_fournisseur = self.mode_button(toggle_row, "Commandes fournisseur (stock)",
                                                 lambda: self.set_type("fournisseur"))
        self.btn_fournisseur.pack(side="left")

        self.search_var = self.make_search_box("Rechercher une commande…")
        self.tree = self.make_table(["Client / Fournisseur", "Description", "Commandée le",
                                      "Prévue le", "Statut", "Alerte"])
        self.tree.bind("<Double-1>", lambda e: self.open_form(self._selected_id()))
        btns = tk.Frame(self, bg=CONTENT_BG)
        btns.pack(fill="x", pady=(8, 0))
        outline_button(btns, "Modifier la sélection", lambda: self.open_form(self._selected_id()), icon="edit").pack(side="left")
        danger_button(btns, "Supprimer", self.delete_selected).pack(side="left", padx=(0, 12))
        self._update_toggle_style()

    def set_type(self, t):
        self.show_type = t
        self._update_toggle_style()
        self.refresh()

    def _update_toggle_style(self):
        self.set_mode_button_active(self.btn_client, self.show_type == "client")
        self.set_mode_button_active(self.btn_fournisseur, self.show_type == "fournisseur")

    def _selected_id(self):
        sel = self.tree.selection()
        return int(sel[0]) if sel else None

    @staticmethod
    def alert_state(commande):
        if commande["statut"] in ("Livrée", "Annulée") or not commande["date_prevue"]:
            return None
        try:
            prevue = datetime.date.fromisoformat(commande["date_prevue"])
        except ValueError:
            return None
        today = datetime.date.today()
        if prevue < today:
            return "overdue"
        if (prevue - today).days <= CommandesPage.ALERT_DAYS:
            return "soon"
        return None

    def refresh(self):
        for i in self.tree.get_children():
            self.tree.delete(i)
        q = self.search_var.get()
        conn = get_connection()
        rows = conn.execute("""SELECT commandes.*, clients.nom as client_nom, clients.prenom as client_prenom
                                FROM commandes LEFT JOIN clients ON clients.id = commandes.client_id
                                WHERE COALESCE(commandes.type,'client')=?
                                ORDER BY date_prevue""", (self.show_type,)).fetchall()
        conn.close()
        n = 0
        for cmd in rows:
            if self.show_type == "client":
                who = f"{cmd['client_nom'] or ''} {cmd['client_prenom'] or ''}".strip() or "—"
            else:
                who = cmd["fournisseur"] or "—"
            if not self.row_matches(q, who, cmd["description"], cmd["statut"]):
                continue
            alert = self.alert_state(cmd)
            alert_txt = {"overdue": "EN RETARD", "soon": "BIENTÔT"}.get(alert, "—")
            tag = "low" if alert else ("even" if n % 2 == 0 else "odd")
            self.tree.insert("", "end", iid=str(cmd["id"]), tags=(tag,),
                              values=(who, cmd["description"] or "—", cmd["date_commande"] or "—",
                                      cmd["date_prevue"] or "—", cmd["statut"], alert_txt))
            n += 1

    def delete_selected(self):
        cid = self._selected_id()
        if not cid:
            messagebox.showinfo("Info", "Sélectionnez une commande.")
            return
        if not messagebox.askyesno("Confirmer", "Supprimer cette commande ?"):
            return
        conn = get_connection()
        conn.execute("DELETE FROM commandes WHERE id=?", (cid,))
        conn.commit(); conn.close()
        log_action(self.app.current_user["username"], "Commande supprimée")
        self.refresh()

    def open_form(self, commande_id=None):
        conn = get_connection()
        cmd = conn.execute("SELECT * FROM commandes WHERE id=?", (commande_id,)).fetchone() if commande_id else None
        clients = conn.execute("SELECT * FROM clients ORDER BY nom").fetchall()
        conn.close()
        client_map = {f"{c['nom']} {c['prenom'] or ''}".strip(): c["id"] for c in clients}
        current_name = ""
        if cmd and cmd["client_id"]:
            for name, cid in client_map.items():
                if cid == cmd["client_id"]:
                    current_name = name

        modal = ModalForm(self.app, "Modifier commande" if cmd else "Nouvelle commande", width=420)
        body = modal.body

        type_v = dropdown(body, "Type de commande", ["client", "fournisseur"],
                           (cmd["type"] if cmd and cmd["type"] else self.show_type))

        client_wrap = tk.Frame(body, bg=SURFACE)
        tk.Label(client_wrap, text="CLIENT", font=FONT_MONO_SM, fg=SLATE, bg=SURFACE, anchor="w").pack(fill="x", pady=(8, 2))
        client_var = tk.StringVar(value=current_name)
        cb = ttk.Combobox(client_wrap, textvariable=client_var, values=list(client_map.keys()),
                           state="readonly", font=FONT_BODY, style="Themed.TCombobox")
        cb.pack(fill="x", ipady=2)

        fournisseur_wrap = tk.Frame(body, bg=SURFACE)
        fournisseur_e = field(fournisseur_wrap, "Fournisseur", cmd["fournisseur"] if cmd else "")

        def sync_type_fields(*_a):
            if type_v.get() == "client":
                fournisseur_wrap.pack_forget()
                client_wrap.pack(fill="x")
            else:
                client_wrap.pack_forget()
                fournisseur_wrap.pack(fill="x")
        type_v.trace_add("write", sync_type_fields)
        sync_type_fields()

        desc_e = field(body, "Description (ex. montures, verres…)", cmd["description"] if cmd else "")
        row = tk.Frame(body, bg=SURFACE); row.pack(fill="x")
        row.grid_columnconfigure(0, weight=1); row.grid_columnconfigure(1, weight=1)
        cmd_wrap = tk.Frame(row, bg=SURFACE); cmd_wrap.grid(row=0, column=0, sticky="ew", padx=(0, 6))
        prevue_wrap = tk.Frame(row, bg=SURFACE); prevue_wrap.grid(row=0, column=1, sticky="ew", padx=(6, 0))
        date_cmd_e = field_date(cmd_wrap, "Commandée le", cmd["date_commande"] if cmd else today_iso())
        date_prevue_e = field_date(prevue_wrap, "Prévue le", cmd["date_prevue"] if cmd else "")
        statut_v = dropdown(body, "Statut", STATUTS_COMMANDE, cmd["statut"] if cmd else STATUTS_COMMANDE[0])

        def save():
            is_client = type_v.get() == "client"
            if is_client and not client_var.get():
                messagebox.showwarning("Champ requis", "Sélectionnez un client.")
                return
            if not is_client and not fournisseur_e.get().strip():
                messagebox.showwarning("Champ requis", "Indiquez le nom du fournisseur.")
                return
            cid = client_map[client_var.get()] if is_client else None
            fournisseur_val = fournisseur_e.get().strip() if not is_client else None
            who_label = client_var.get() if is_client else fournisseur_val
            conn = get_connection()
            if cmd:
                conn.execute("""UPDATE commandes SET client_id=?, fournisseur=?, type=?, description=?,
                                 date_commande=?, date_prevue=?, statut=? WHERE id=?""",
                             (cid, fournisseur_val, type_v.get(), desc_e.get(), date_cmd_e.get(),
                              date_prevue_e.get() or None, statut_v.get(), cmd["id"]))
                action = f"Commande modifiée : {who_label}"
            else:
                conn.execute("""INSERT INTO commandes (client_id,fournisseur,type,description,date_commande,
                                 date_prevue,statut) VALUES (?,?,?,?,?,?,?)""",
                             (cid, fournisseur_val, type_v.get(), desc_e.get(), date_cmd_e.get(),
                              date_prevue_e.get() or None, statut_v.get()))
                action = f"Commande créée : {who_label}"
            conn.commit(); conn.close()
            log_action(self.app.current_user["username"], action)
            modal.destroy()
            self.set_type(type_v.get())

        modal.buttons(save)



# ================================================================== ENTRY POINT
if __name__ == "__main__":
    init_db()
    app = OptiluxApp()
    app.mainloop()
